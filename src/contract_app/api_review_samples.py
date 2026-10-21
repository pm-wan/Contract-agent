"""Downloadable source and contract fixtures run through the real review workflow."""

from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi import APIRouter, Depends, HTTPException
from starlette.responses import Response

from contract_app.auth_security import require_permission
from contract_app.rules.sales_contract_review import list_scenarios

router = APIRouter(
    prefix="/api/v1/review-samples",
    tags=["Review samples"],
    dependencies=[Depends(require_permission("cases.use"))],
)

SOURCE_SAMPLES = [
    {
        "key": "civil-code",
        "kind": "legal",
        "title": "中华人民共和国民法典（全文）",
        "issuer": "全国人民代表大会",
        "effective_from": "2021-01-01",
        "scope": "中华人民共和国",
        "source_url": "https://www.beijing.gov.cn/zhengce/zhengcefagui/202006/t20200602_1913330.html",
    },
    {
        "key": "sales-policy",
        "kind": "company",
        "title": "测试资料：销售管理制度（2026版）",
        "issuer": "演示公司",
        "effective_from": "2026-01-01",
        "scope": "演示公司作为卖方的销售合同",
        "source_url": "",
        "content": """销售管理制度（测试资料，不是本公司的真实制度）
第一条 本制度适用于演示公司作为卖方签订的销售合同，自2026年1月1日起生效。
第二条 销售折扣不得超过标准价的15%；超过15%时须取得销售总监书面批准。
第三条 客户账期不得超过60天；超过60天时须取得财务负责人书面批准。
第四条 合同含税总额达到100万元（含）时，须取得总经理书面审批。
第五条 偏离公司标准销售合同模板的责任、验收或付款约定时，须经法务书面复核。
第六条 审批文件应随合同归档。未提交审批文件时应要求补充核实。
""",
    },
    {
        "key": "old-sales-policy",
        "kind": "company",
        "title": "测试资料：销售管理制度（历史版）",
        "issuer": "演示公司",
        "effective_from": "2024-01-01",
        "effective_until": "2025-12-31",
        "scope": "演示公司作为卖方的销售合同",
        "source_url": "",
        "content": "历史销售管理制度（测试资料）\n本版本适用于2024年1月1日至2025年12月31日的合同。"
        "折扣超过10%须销售总监书面批准；账期超过30天须财务负责人书面批准。",
    },
]


@router.get("")
def list_samples() -> dict:
    return {
        "sources": [{k: v for k, v in item.items() if k != "content"} for item in SOURCE_SAMPLES],
        "contracts": [
            {
                "scenario_id": item["scenario_id"],
                "title": item["title"],
                "description": item["description"],
                "category": item["category"],
            }
            for item in list_scenarios()
        ],
    }


@router.get("/sources/{key}")
def download_sample_source(key: str) -> Response:
    item = next((source for source in SOURCE_SAMPLES if source["key"] == key), None)
    if item is None:
        raise HTTPException(404, "sample_not_found")
    if key == "civil-code":
        data = (Path(__file__).parent / "sample_data" / "civil_code.txt").read_bytes()
    else:
        data = item["content"].encode("utf-8")
    return Response(data, media_type="text/plain; charset=utf-8")


@router.get("/contracts/{scenario_id}")
def download_sample_contract(scenario_id: str) -> Response:
    item = next(
        (source for source in list_scenarios() if source["scenario_id"] == scenario_id), None
    )
    if item is None:
        raise HTTPException(404, "sample_not_found")
    document = Document()
    document.add_heading("销售合同 · " + item["title"], 0)
    document.add_paragraph("测试合同：用于验证合同理解、资料检索、引用和比对流程。")
    facts = item["facts"]
    document.add_paragraph("甲方（买方）：演示采购有限公司；乙方（卖方、我方）：演示公司。")
    document.add_paragraph(
        "本合同适用中华人民共和国法律。签订日期：" + str(facts.get("contract_date", "待确认"))
    )
    document.add_paragraph(item["contract_text"])
    if facts.get("amount") is not None:
        document.add_paragraph(f"合同含税总价为人民币{facts['amount']}元。")
    if facts.get("discount_percent") is not None:
        document.add_paragraph(f"成交价格在公司标准价基础上优惠{facts['discount_percent']}%。")
    if facts.get("credit_days") is not None:
        document.add_paragraph(f"买方应于验收后{facts['credit_days']}日内支付剩余款项。")
    if facts.get("template_deviation"):
        document.add_paragraph("本合同的责任与验收条款已对公司标准模板作修改。")
    if facts.get("approval_obtained") is True:
        document.add_paragraph(
            "附件：销售总监、财务负责人及总经理已书面批准本次折扣、账期及合同金额。"
        )
    elif facts.get("approval_obtained") is False:
        document.add_paragraph("双方确认本合同尚未取得公司内部授权审批。")
    if facts.get("template_approval") is True:
        document.add_paragraph("附件：法务已书面复核本合同模板偏离。")
    document.add_paragraph("双方签章：买方代表________；卖方代表________。")
    buffer = BytesIO()
    document.save(buffer)
    return Response(
        buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
