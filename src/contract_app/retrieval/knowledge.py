"""Source-document parsing, immutable snapshots and query-driven passage retrieval."""

from __future__ import annotations

import json
import re
from datetime import date
from hashlib import sha256
from io import BytesIO
from typing import Any
from uuid import UUID

import psycopg
from psycopg.rows import dict_row

from contract_app.auth_security import set_tenant

MAX_SOURCE_BYTES = 10 * 1024 * 1024
MAX_SNAPSHOT_BYTES = 4 * 1024 * 1024


def digest(value: Any) -> str:
    return sha256(
        json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()


def parse_source(data: bytes, file_name: str) -> tuple[str, list[dict[str, str]], list[str]]:
    if not data or len(data) > MAX_SOURCE_BYTES:
        raise ValueError("knowledge_file_size_invalid")
    suffix = file_name.rsplit(".", 1)[-1].lower()
    warnings: list[str] = []
    if suffix in {"txt", "md"}:
        text = data.decode("utf-8-sig")
        segments = [
            {"locator": f"段落 {i}", "text": text.strip()}
            for i, text in enumerate(re.split(r"\n\s*\n|\r?\n", text), 1)
            if text.strip()
        ]
        grouped = []
        article = None
        for segment in segments:
            match = re.match(r"^(第[零〇一二三四五六七八九十百千万\d]+条)", segment["text"])
            heading = re.match(r"^第[零〇一二三四五六七八九十百千万\d]+[编章节]", segment["text"])
            if match or heading:
                article = {"locator": match.group(1), "text": segment["text"]} if match else None
                grouped.append(article if article is not None else segment)
            elif article is not None:
                article["text"] += "\n" + segment["text"]
            else:
                grouped.append(segment)
        segments = grouped
        media_type = "text/plain"
    elif suffix == "docx":
        from docx import Document

        from contract_app.documents.parser import _check_input

        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        _check_input(data, media_type)
        document = Document(BytesIO(data))
        texts = [p.text for p in document.paragraphs]
        texts += [
            " | ".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        ]
        segments = [
            {"locator": f"段落/表格 {i}", "text": text.strip()}
            for i, text in enumerate(texts, 1)
            if text.strip()
        ]
        warnings = ["DOCX 按段落及表格定位；上传前请确认正文覆盖页眉、文本框等特殊内容。"]
    elif suffix == "pdf":
        from contract_app.documents.parser import parse_document

        media_type = "application/pdf"
        parsed = parse_document(data, media_type)
        segments = [
            {"locator": f"第 {i} 页", "text": text}
            for i, text in enumerate(parsed.pages, 1)
            if text.strip()
        ]
        warnings = list(parsed.warnings)
    else:
        raise ValueError("knowledge_file_type_unsupported")
    # Split long passages for bounded tool output, preserving source anchors.
    chunks = []
    for segment in segments:
        for offset in range(0, len(segment["text"]), 2400):
            chunks.append(
                {
                    "locator": segment["locator"] + (f" · 字符 {offset}" if offset else ""),
                    "text": segment["text"][offset : offset + 2400],
                }
            )
    if not chunks or len(chunks) > 20000:
        raise ValueError("knowledge_text_missing_or_too_large")
    return media_type, chunks, warnings


def serialize_source(row: dict[str, Any]) -> dict[str, Any]:
    return {
        key: str(value) if isinstance(value, (UUID, date)) else value
        for key, value in row.items()
        if key != "file_data"
    }


def capture_sources(database_url: str, tenant_id: UUID) -> dict[str, Any]:
    with psycopg.connect(database_url, row_factory=dict_row) as connection:
        set_tenant(connection, tenant_id)
        rows = connection.execute(
            """SELECT id,source_key,version,kind,title,source_url,issuer,
            scope,effective_from,effective_until,content_sha256,segments,warnings
            FROM knowledge_documents WHERE tenant_id=%s AND status='published'
            ORDER BY source_key,version DESC""",
            (tenant_id,),
        ).fetchall()
    # Keep all published versions: applicability depends on the contract date.
    sources = [serialize_source(row) for row in rows]
    if len(json.dumps(sources, ensure_ascii=False).encode()) > MAX_SNAPSHOT_BYTES:
        raise ValueError("knowledge_snapshot_too_large")
    return {"sources": sources, "sources_sha256": digest(sources)}


def search_sources(
    sources: list[dict[str, Any]], kind: str, query: str, keywords: list[str]
) -> list[dict[str, Any]]:
    """Model-selected terms rank actual passages; no check IDs or legal rules exist here."""
    terms = list(dict.fromkeys(term.strip().casefold() for term in keywords if term.strip()))
    if not terms:
        terms = re.findall(r"[\w\u4e00-\u9fff]{2,}", query.casefold())
    ranked = []
    for source in sources:
        if source["kind"] != kind:
            continue
        for index, segment in enumerate(source["segments"]):
            text = segment["text"].casefold()
            score = sum(min(5, text.count(term)) * max(1, len(term)) for term in terms)
            if score:
                ranked.append((score, str(source["id"]), index, source, segment))
    ranked.sort(key=lambda item: (-item[0], item[1], item[2]))
    return [
        {
            "source_id": f"local:{source['id']}:{index}",
            "kind": kind,
            "origin": "local",
            "document_id": source["id"],
            "title": source["title"],
            "url": source["source_url"],
            "issuer": source["issuer"],
            "scope": source["scope"],
            "version": source["version"],
            "content_sha256": source["content_sha256"],
            "effective_from": source["effective_from"],
            "effective_until": source["effective_until"],
            "locator": segment["locator"],
            "text": segment["text"],
        }
        for _, _, index, source, segment in ranked[:8]
    ]
