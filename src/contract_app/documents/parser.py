"""Extract stable Unicode text and evidence locations from PDF and DOCX."""

from dataclasses import replace
from hashlib import sha256
from io import BytesIO
from typing import Any, Literal, cast
from zipfile import BadZipFile, ZipFile

import pymupdf
from docx import Document

from contract_app.documents.docx_render import RenderedDocx, render_docx_to_pdf
from contract_app.documents.models import DOCX_MEDIA_TYPE as DOCX_MEDIA_TYPE
from contract_app.documents.models import PDF_MEDIA_TYPE as PDF_MEDIA_TYPE
from contract_app.documents.models import ParsedDocument, TextBlock
from contract_app.documents.ocr import OcrEngine, OcrError, OcrWord

fitz = cast(Any, pymupdf)

MAX_DOCUMENT_BYTES = 30 * 1024 * 1024
MAX_DOCX_UNCOMPRESSED_BYTES = 120 * 1024 * 1024
PARSER_VERSION = "m0-pymupdf-docx-1"
DOCX_PARSER_VERSION = "m1-pymupdf-docx-render-2"
OCR_RENDER_SCALE = 2.0
MAX_OCR_PIXELS = 20_000_000


class ParseError(ValueError):
    """A document cannot be parsed safely or completely enough for review."""


def _check_input(data: bytes, media_type: str) -> None:
    if not data or len(data) > MAX_DOCUMENT_BYTES:
        raise ParseError("document_size_invalid")
    if media_type == PDF_MEDIA_TYPE and not data.startswith(b"%PDF-"):
        raise ParseError("pdf_signature_invalid")
    if media_type == DOCX_MEDIA_TYPE:
        if not data.startswith(b"PK\x03\x04"):
            raise ParseError("docx_signature_invalid")
        try:
            with ZipFile(BytesIO(data)) as archive:
                names = set(archive.namelist())
                if "word/document.xml" not in names:
                    raise ParseError("docx_structure_invalid")
                total_size = sum(info.file_size for info in archive.infolist())
                if total_size > MAX_DOCX_UNCOMPRESSED_BYTES:
                    raise ParseError("docx_expanded_size_invalid")
                if any(name.endswith("vbaProject.bin") for name in names):
                    raise ParseError("docx_macro_forbidden")
        except BadZipFile as exc:
            raise ParseError("docx_archive_invalid") from exc
    elif media_type != PDF_MEDIA_TYPE:
        raise ParseError("media_type_unsupported")


def parse_document(
    data: bytes,
    media_type: str,
    *,
    ocr: OcrEngine | None = None,
    rendered_docx: RenderedDocx | None = None,
) -> ParsedDocument:
    """Parse a file without mutating the original byte stream.

    Offsets are Python Unicode code-point indexes into each normalized page string.
    DOCX page locations refer to the versioned canonical PDF rendering, not Word's
    environment-dependent pagination.
    """
    _check_input(data, media_type)
    digest = sha256(data).hexdigest()
    if media_type == PDF_MEDIA_TYPE:
        return replace(_parse_pdf(data, digest, ocr), index_source_sha256=digest)
    rendered = rendered_docx or render_docx_to_pdf(data)
    if rendered.source_sha256 != digest:
        raise ParseError("docx_render_source_hash_mismatch")
    rendered_digest = sha256(rendered.pdf_bytes).hexdigest()
    pdf = _parse_pdf(rendered.pdf_bytes, rendered_digest, None)
    return ParsedDocument(
        source_sha256=digest,
        media_type=DOCX_MEDIA_TYPE,
        parser_version=DOCX_PARSER_VERSION,
        pages=pdf.pages,
        blocks=pdf.blocks,
        warnings=tuple(dict.fromkeys((*pdf.warnings, "docx_pages_use_canonical_pdf_rendering"))),
        index_source_sha256=rendered_digest,
        renderer_identity=rendered.renderer_identity,
    )


def _parse_pdf(data: bytes, digest: str, ocr: OcrEngine | None) -> ParsedDocument:
    pages: list[str] = []
    blocks: list[TextBlock] = []
    warnings: list[str] = []
    ocr_used = False
    try:
        with fitz.open(stream=data, filetype="pdf") as pdf:
            if pdf.is_encrypted:
                raise ParseError("pdf_encrypted")
            for page_index, page in enumerate(pdf):
                page_chunks: list[str] = []
                offset = 0
                for item in page.get_text("blocks", sort=True):
                    if item[6] != 0:
                        continue
                    text = str(item[4]).strip()
                    if not text:
                        continue
                    if page_chunks:
                        page_chunks.append("\n")
                        offset += 1
                    start = offset
                    page_chunks.append(text)
                    offset += len(text)
                    blocks.append(
                        TextBlock(
                            page=page_index + 1,
                            text=text,
                            start_offset=start,
                            end_offset=offset,
                            bbox=(float(item[0]), float(item[1]), float(item[2]), float(item[3])),
                        )
                    )
                page_text = "".join(page_chunks)
                if not page_text:
                    if ocr is None:
                        warnings.append(f"page_{page_index + 1}_has_no_text_layer")
                    else:
                        ocr_used = True
                        page_text, ocr_blocks, ocr_warnings = _ocr_page(page, page_index + 1, ocr)
                        blocks.extend(ocr_blocks)
                        warnings.extend(ocr_warnings)
                elif page.get_images(full=True):
                    warnings.append(f"page_{page_index + 1}_has_images_unverified")
                pages.append(page_text)
    except ParseError:
        raise
    except OcrError as exc:
        code = str(exc)
        if code not in {
            "ocr_image_invalid",
            "ocr_unavailable",
            "ocr_timeout",
            "ocr_failed",
            "ocr_output_too_large",
            "ocr_output_invalid",
        }:
            code = "ocr_failed"
        raise ParseError(code) from exc
    except Exception as exc:
        raise ParseError("pdf_parse_failed") from exc
    if not pages:
        raise ParseError("pdf_has_no_pages")
    version = f"{PARSER_VERSION}+{ocr.version}" if ocr_used and ocr is not None else PARSER_VERSION
    return ParsedDocument(
        digest,
        PDF_MEDIA_TYPE,
        version,
        tuple(pages),
        tuple(blocks),
        tuple(warnings),
        index_source_sha256=digest,
    )


def _ocr_page(
    page: Any, page_number: int, engine: OcrEngine
) -> tuple[str, list[TextBlock], list[str]]:
    pixel_width = round(page.rect.width * OCR_RENDER_SCALE)
    pixel_height = round(page.rect.height * OCR_RENDER_SCALE)
    if pixel_width <= 0 or pixel_height <= 0 or pixel_width * pixel_height > MAX_OCR_PIXELS:
        raise ParseError("ocr_page_too_large")
    pixmap = page.get_pixmap(matrix=fitz.Matrix(OCR_RENDER_SCALE, OCR_RENDER_SCALE), alpha=False)
    words = engine.recognize_png(pixmap.tobytes("png"))
    warnings = [f"page_{page_number}_ocr_requires_review"]
    if not words:
        warnings.append(f"page_{page_number}_ocr_no_text")
        return "", [], warnings
    groups: dict[tuple[int, int, int], list[OcrWord]] = {}
    for word in words:
        if (
            word.left < 0
            or word.top < 0
            or word.left + word.width > pixmap.width
            or word.top + word.height > pixmap.height
            or not word.text.strip()
            or not 0 <= word.confidence <= 100
        ):
            raise ParseError("ocr_coordinates_invalid")
        groups.setdefault(word.line_id, []).append(word)
    page_parts: list[str] = []
    blocks: list[TextBlock] = []
    offset = 0
    scale_x = page.rect.width / pixmap.width
    scale_y = page.rect.height / pixmap.height
    for line_words in groups.values():
        text = " ".join(word.text for word in line_words)
        if page_parts:
            page_parts.append("\n")
            offset += 1
        start = offset
        page_parts.append(text)
        offset += len(text)
        left = min(word.left for word in line_words)
        top = min(word.top for word in line_words)
        right = max(word.left + word.width for word in line_words)
        bottom = max(word.top + word.height for word in line_words)
        confidence = min(word.confidence for word in line_words) / 100
        blocks.append(
            TextBlock(
                page=page_number,
                text=text,
                start_offset=start,
                end_offset=offset,
                bbox=(
                    page.rect.x0 + left * scale_x,
                    page.rect.y0 + top * scale_y,
                    page.rect.x0 + right * scale_x,
                    page.rect.y0 + bottom * scale_y,
                ),
                kind="ocr",
                confidence=confidence,
            )
        )
        if confidence < 0.8:
            warnings.append(f"page_{page_number}_ocr_low_confidence")
    return "".join(page_parts), blocks, list(dict.fromkeys(warnings))


def _docx_text(element: Any) -> str:
    chunks: list[str] = []
    for node in element.iter():
        if node.tag.endswith("}t"):
            chunks.append(node.text or "")
        elif node.tag.endswith("}tab"):
            chunks.append("\t")
        elif node.tag.endswith("}br") or node.tag.endswith("}cr"):
            chunks.append("\n")
    return "".join(chunks)


def _parse_docx(data: bytes, digest: str) -> ParsedDocument:
    chunks: list[str] = []
    blocks: list[TextBlock] = []
    warnings = ["docx_has_no_stable_page_coordinates"]
    offset = 0
    try:
        document = Document(BytesIO(data))
        header_attributes = (
            "header",
            "first_page_header",
            "even_page_header",
        )
        footer_attributes = (
            "footer",
            "first_page_footer",
            "even_page_footer",
        )
        if any(
            _docx_text(getattr(section, attribute)._element).strip()
            for section in document.sections
            for attribute in header_attributes
        ):
            warnings.append("docx_header_content_not_reviewed")
        if any(
            _docx_text(getattr(section, attribute)._element).strip()
            for section in document.sections
            for attribute in footer_attributes
        ):
            warnings.append("docx_footer_content_not_reviewed")
        for element in document.element.body:
            kind: Literal["paragraph", "table"]
            if element.tag.endswith("}p"):
                text = _docx_text(element)
                kind = "paragraph"
            elif element.tag.endswith("}tbl"):
                rows = []
                for row in element:
                    if not row.tag.endswith("}tr"):
                        continue
                    cells = []
                    for cell in row:
                        if cell.tag.endswith("}tc"):
                            cells.append(_docx_text(cell))
                    rows.append("\t".join(cells))
                text = "\n".join(rows)
                kind = "table"
            else:
                continue
            text = text.strip()
            if not text:
                continue
            if chunks:
                chunks.append("\n")
                offset += 1
            start = offset
            chunks.append(text)
            offset += len(text)
            blocks.append(TextBlock(1, text, start, offset, None, kind))
    except Exception as exc:
        raise ParseError("docx_parse_failed") from exc
    page_text = "".join(chunks)
    if not page_text:
        warnings.append("docx_has_no_text")
    return ParsedDocument(
        digest, DOCX_MEDIA_TYPE, DOCX_PARSER_VERSION, (page_text,), tuple(blocks), tuple(warnings)
    )
