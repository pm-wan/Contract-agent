import hashlib
import json
import os
import shutil
from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path
from threading import Barrier
from typing import Any
from uuid import UUID, uuid4

import fitz
import psycopg
import pytest
from docx import Document
from fastapi.testclient import TestClient

from contract_app.api import create_app
from contract_app.api_documents import (
    get_document_authorization_store,
    get_document_catalog_store,
    get_document_content_store,
    get_document_ingestion_store,
)
from contract_app.api_human_tasks import get_job_store as get_human_task_job_store
from contract_app.api_runs import get_job_store
from contract_app.api_snapshots import Principal, get_snapshot_store, require_principal
from contract_app.cases.legal_entity import CaseLegalEntityStore
from contract_app.contextassembly.snapshot import InputSnapshotStore, PreviewConflict
from contract_app.documents.authorization import DocumentAuthorizationStore
from contract_app.documents.catalog import DocumentCatalogStore
from contract_app.documents.content import DocumentContentStore
from contract_app.documents.docx_render import DocxRenderError, RenderedDocx
from contract_app.documents.index_jobs import DocumentIndexJobStore
from contract_app.documents.index_worker import DocumentIndexWorker
from contract_app.documents.ingestion import DocumentIngestionStore
from contract_app.documents.object_gc import DocumentObjectGarbageCollector
from contract_app.documents.parse_jobs import DocumentParseJobStore
from contract_app.documents.parse_worker import DocumentParseWorker
from contract_app.documents.parser import DOCX_MEDIA_TYPE, PDF_MEDIA_TYPE
from contract_app.documents.storage import LocalObjectStore, ObjectRef
from contract_app.retrieval.authorization import AuthorizationProjector, TaskAccess
from contract_app.retrieval.pageindex_adapter import PageIndexAdapter, RetrievalDenied
from contract_app.workflow.execution_trace import ExecutionTraceRecorder
from contract_app.workflow.job_store import BoundRunConflict, JobStore

TEST_DATABASE_URL = os.environ.get("TEST_DATABASE_URL")
TEST_ADMIN_DATABASE_URL = os.environ.get("TEST_ADMIN_DATABASE_URL")
pytestmark = pytest.mark.skipif(
    not TEST_DATABASE_URL or not TEST_ADMIN_DATABASE_URL,
    reason="both TEST_DATABASE_URL and TEST_ADMIN_DATABASE_URL are required",
)


def _case() -> tuple[UUID, UUID, UUID]:
    assert TEST_ADMIN_DATABASE_URL is not None
    tenant, actor, entity = uuid4(), uuid4(), uuid4()
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        case = connection.execute(
            "INSERT INTO audit_cases (tenant_id, title, legal_entity_id) "
            "VALUES (%s, 'upload fixture', %s) RETURNING id",
            (tenant, entity),
        ).fetchone()
        assert case is not None
        connection.execute(
            "INSERT INTO case_members (tenant_id, case_id, actor_id, role) "
            "VALUES (%s, %s, %s, 'reviewer')",
            (tenant, case[0], actor),
        )
    return tenant, case[0], actor


def _pdf(text: str = "Payment is due within 30 days.") -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    data = document.tobytes()
    document.close()
    return data


def _docx(text: str = "Payment is due within 30 days.") -> bytes:
    document = Document()
    document.add_paragraph(text)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _render_docx_for_test(data: bytes) -> RenderedDocx:
    document = Document(BytesIO(data))
    pdf = fitz.open()
    page = pdf.new_page()
    content = "\n".join(paragraph.text for paragraph in document.paragraphs)
    page.insert_textbox(fitz.Rect(48, 48, 550, 790), content, fontsize=12)
    return RenderedDocx(pdf.tobytes(), "test-renderer:1", hashlib.sha256(data).hexdigest())


def _uploaded_count(tenant: UUID, case_id: UUID) -> int:
    assert TEST_DATABASE_URL is not None
    with psycopg.connect(TEST_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        return connection.execute(
            "SELECT count(*) FROM document_versions WHERE uploaded_case_id = %s", (case_id,)
        ).fetchone()[0]


def _client(tenant: UUID, actor: UUID, object_store: LocalObjectStore) -> TestClient:
    assert TEST_DATABASE_URL is not None
    app = create_app()
    app.dependency_overrides[require_principal] = lambda: Principal(tenant, actor)
    app.dependency_overrides[get_document_ingestion_store] = lambda: DocumentIngestionStore(
        TEST_DATABASE_URL, object_store
    )
    app.dependency_overrides[get_document_authorization_store] = lambda: DocumentAuthorizationStore(
        TEST_DATABASE_URL
    )
    app.dependency_overrides[get_document_catalog_store] = lambda: DocumentCatalogStore(
        TEST_DATABASE_URL
    )
    app.dependency_overrides[get_document_content_store] = lambda: DocumentContentStore(
        TEST_DATABASE_URL, object_store
    )
    app.dependency_overrides[get_job_store] = lambda: JobStore(TEST_DATABASE_URL)
    app.dependency_overrides[get_human_task_job_store] = lambda: JobStore(TEST_DATABASE_URL)
    app.dependency_overrides[get_snapshot_store] = lambda: InputSnapshotStore(TEST_DATABASE_URL)
    return TestClient(app)


def _parse_once(
    tenant: UUID,
    object_store: LocalObjectStore,
    *,
    use_host_renderer: bool = False,
) -> str:
    assert TEST_DATABASE_URL is not None
    renderer = {} if use_host_renderer else {"docx_renderer": _render_docx_for_test}
    return DocumentParseWorker(
        DocumentParseJobStore(TEST_DATABASE_URL),
        tenant,
        "test-parser",
        object_store,
        lease_seconds=30,
        **renderer,
    ).run_once()


def test_upload_stages_original_and_parse_without_grant_or_selection(tmp_path: Path) -> None:
    assert TEST_DATABASE_URL is not None
    tenant, case_id, actor = _case()
    store = LocalObjectStore(tmp_path)
    client = _client(tenant, actor, store)
    data = _pdf()
    response = client.post(
        f"/api/v1/cases/{case_id}/documents",
        headers={"Idempotency-Key": "upload-00000001"},
        data={"role": "primary_contract"},
        files={"file": ("../signed-contract.pdf", data, PDF_MEDIA_TYPE)},
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["file_name"] == "signed-contract.pdf"
    assert body["document_status"] == "staged"
    assert body["authorization_status"] == "not_granted_by_upload"
    assert body["parse_status"] == "queued"
    assert body["parsed_hash"] is None
    assert body["case_version"] == 2
    assert _uploaded_count(tenant, case_id) == 1

    with psycopg.connect(TEST_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        row = connection.execute(
            "SELECT status, source_kind, has_restricted_amount, upload_case_version, "
            "source_sha256, parsed_sha256, size_bytes, parsed_size_bytes "
            "FROM document_versions WHERE id = %s",
            (body["document_version_id"],),
        ).fetchone()
        assert row is not None
        assert row[:4] == ("staged", "main", True, 2)
        assert row[5] is None and row[7] is None
        assert (
            connection.execute(
                "SELECT count(*) FROM document_parse_jobs WHERE tenant_id = %s AND case_id = %s",
                (tenant, case_id),
            ).fetchone()[0]
            == 1
        )
        assert (
            connection.execute(
                "SELECT count(*) FROM contract_selections WHERE tenant_id = %s AND case_id = %s",
                (tenant, case_id),
            ).fetchone()[0]
            == 0
        )
        assert (
            connection.execute(
                "SELECT count(*) FROM document_grants WHERE tenant_id = %s",
                (tenant,),
            ).fetchone()[0]
            == 0
        )
    original_ref = ObjectRef(tenant, case_id, row[4], row[6])
    assert store.get(original_ref) == data
    content_route = f"/api/v1/cases/{case_id}/documents/{body['document_version_id']}/content"
    original_response = client.get(content_route)
    assert original_response.status_code == 200
    assert original_response.content == data
    assert original_response.headers["cache-control"] == "no-store"
    parsed_response = client.get(content_route, params={"kind": "parsed"})
    assert parsed_response.status_code == 409
    assert _parse_once(tenant, store) == "completed"
    parsed_response = client.get(content_route, params={"kind": "parsed"})
    assert parsed_response.status_code == 200
    assert parsed_response.headers["content-type"] == "application/json"
    assert "30 days" in json.loads(parsed_response.content)["pages"][0]
    client.app.dependency_overrides[require_principal] = lambda: Principal(tenant, uuid4())
    hidden = client.get(content_route)
    assert hidden.status_code == 404
    assert data not in hidden.content


def test_upload_is_idempotent_and_conflicting_replay_is_rejected(tmp_path: Path) -> None:
    tenant, case_id, actor = _case()
    store = LocalObjectStore(tmp_path)
    client = _client(tenant, actor, store)
    data = _pdf()
    route = f"/api/v1/cases/{case_id}/documents"
    request = {
        "headers": {"Idempotency-Key": "upload-00000002"},
        "data": {"role": "attachment"},
        "files": {"file": ("appendix.pdf", data, PDF_MEDIA_TYPE)},
    }
    first = client.post(route, **request)
    replay = client.post(route, **request)
    assert first.status_code == replay.status_code == 201
    assert replay.json() == first.json()
    assert replay.json()["role"] == "attachment"
    assert replay.json()["case_version"] == 2
    assert _uploaded_count(tenant, case_id) == 1
    with psycopg.connect(TEST_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        assert connection.execute(
            "SELECT count(*) FROM document_parse_jobs WHERE tenant_id = %s AND case_id = %s",
            (tenant, case_id),
        ).fetchone() == (1,)
    conflict = client.post(
        route,
        headers={"Idempotency-Key": "upload-00000002"},
        data={"role": "attachment"},
        files={"file": ("appendix.pdf", _pdf("Different bytes."), PDF_MEDIA_TYPE)},
    )
    assert conflict.status_code == 409
    assert "idempotency_key_conflict" not in conflict.text
    assert _uploaded_count(tenant, case_id) == 1


def test_document_parse_claim_reports_queue_wait(tmp_path: Path) -> None:
    assert TEST_DATABASE_URL is not None and TEST_ADMIN_DATABASE_URL is not None
    tenant, case_id, actor = _case()
    client = _client(tenant, actor, LocalObjectStore(tmp_path))
    uploaded = client.post(
        f"/api/v1/cases/{case_id}/documents",
        headers={"Idempotency-Key": "parse-wait-upload-0001"},
        data={"role": "primary_contract"},
        files={"file": ("parse-wait.pdf", _pdf(), PDF_MEDIA_TYPE)},
    )
    assert uploaded.status_code == 201, uploaded.text
    document_id = UUID(uploaded.json()["document_version_id"])
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        connection.execute(
            "UPDATE document_parse_jobs SET created_at = now() - interval '2 seconds', "
            "next_run_at = now() - interval '2 seconds' "
            "WHERE document_version_id = %s",
            (document_id,),
        )

    store = DocumentParseJobStore(TEST_DATABASE_URL)
    job = store.claim(tenant, "queue-wait-parser")
    assert job is not None and job.document_version_id == document_id
    assert job.queue_wait_ms >= 1500
    assert store.fail(tenant, job, "queue-wait-parser", "telemetry_probe") == "failed"


def test_document_index_claim_reports_queue_wait(tmp_path: Path) -> None:
    assert TEST_DATABASE_URL is not None and TEST_ADMIN_DATABASE_URL is not None
    tenant, case_id, actor = _case()
    object_store = LocalObjectStore(tmp_path)
    client = _client(tenant, actor, object_store)
    uploaded = client.post(
        f"/api/v1/cases/{case_id}/documents",
        headers={"Idempotency-Key": "index-wait-upload-0001"},
        data={"role": "primary_contract"},
        files={"file": ("index-wait.pdf", _pdf(), PDF_MEDIA_TYPE)},
    )
    assert uploaded.status_code == 201, uploaded.text
    assert _parse_once(tenant, object_store) == "completed"
    document_id = UUID(uploaded.json()["document_version_id"])
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        connection.execute(
            "UPDATE document_index_jobs SET created_at = now() - interval '2 seconds', "
            "next_run_at = now() - interval '2 seconds' "
            "WHERE document_version_id = %s",
            (document_id,),
        )

    store = DocumentIndexJobStore(TEST_DATABASE_URL)
    job = store.claim(tenant, "queue-wait-indexer")
    assert job is not None and job.document_version_id == document_id
    assert job.queue_wait_ms >= 1500
    assert (
        store.retry_or_fail(
            tenant,
            job,
            "queue-wait-indexer",
            "telemetry_probe",
            retryable=False,
        )
        == "failed"
    )


def test_case_document_catalog_separates_index_selection_and_actor_authorization(
    tmp_path: Path,
) -> None:
    assert TEST_DATABASE_URL is not None and TEST_ADMIN_DATABASE_URL is not None
    tenant, case_id, actor = _case()
    client = _client(tenant, actor, LocalObjectStore(tmp_path))
    uploaded_response = client.post(
        f"/api/v1/cases/{case_id}/documents",
        headers={"Idempotency-Key": "catalog-upload-0001"},
        data={"role": "primary_contract"},
        files={"file": ("catalog.pdf", _pdf(), PDF_MEDIA_TYPE)},
    )
    assert uploaded_response.status_code == 201, uploaded_response.text
    assert _parse_once(tenant, LocalObjectStore(tmp_path)) == "completed"
    uploaded = uploaded_response.json()
    route = f"/api/v1/cases/{case_id}/documents"

    before = client.get(route)
    assert before.status_code == 200, before.text
    assert len(before.json()) == 1
    assert before.json()[0]["document_status"] == "staged"
    assert before.json()[0]["index_status"] == "queued"
    assert before.json()[0]["selection_id"] is None
    assert before.json()[0]["selection_status"] == "not_selected"
    assert before.json()[0]["authorization_status"] == "not_authorized"
    assert before.json()[0]["authorization_valid_until"] is None

    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        connection.execute(
            "INSERT INTO document_indexes "
            "(tenant_id, document_version_id, source_sha256, pageindex_doc_id) "
            "VALUES (%s, %s, %s, %s)",
            (
                tenant,
                uploaded["document_version_id"],
                uploaded["content_hash"],
                "catalog-pageindex-document",
            ),
        )
    valid_until = datetime.now(UTC) + timedelta(minutes=30)
    selected = client.post(
        f"/api/v1/cases/{case_id}/document-selections",
        json={
            "document_version_id": uploaded["document_version_id"],
            "relation": "main",
            "allow_amount": True,
            "valid_until": valid_until.isoformat(),
        },
    )
    assert selected.status_code == 201, selected.text
    after = client.get(route).json()[0]
    assert after["document_status"] == "active"
    assert after["index_status"] == "ready"
    assert after["selection_id"] == selected.json()["selection_id"]
    assert after["selection_status"] == "selected"
    assert after["authorization_status"] == "authorized"
    assert after["amount_authorized"] is True
    assert after["authorization_valid_until"] == selected.json()["valid_until"]

    client.app.dependency_overrides[require_principal] = lambda: Principal(tenant, uuid4())
    denied = client.get(route)
    assert denied.status_code == 403
    assert uploaded["content_hash"] not in denied.text


def test_document_index_worker_builds_hash_bound_index_from_queued_upload(tmp_path: Path) -> None:
    assert TEST_DATABASE_URL is not None
    tenant, case_id, actor = _case()
    object_store = LocalObjectStore(tmp_path)
    client = _client(tenant, actor, object_store)
    uploaded = client.post(
        f"/api/v1/cases/{case_id}/documents",
        headers={"Idempotency-Key": "index-upload-0001"},
        data={"role": "primary_contract"},
        files={"file": ("index-me.pdf", _pdf("Payment after delivery."), PDF_MEDIA_TYPE)},
    )
    assert uploaded.status_code == 201, uploaded.text
    assert _parse_once(tenant, object_store) == "completed"

    class FakePageIndexClient:
        def submit_document(self, file_path: str, *, wait: bool = False) -> dict[str, Any]:
            assert wait and Path(file_path).read_bytes().startswith(b"%PDF-")
            return {"doc_id": "queued-index-document"}

        def get_tree(
            self, doc_id: str, node_summary: bool = False, include_text: bool = True
        ) -> dict[str, Any]:
            return {"doc_id": doc_id}

        def get_page_content(self, doc_id: str, pages: str) -> list[dict[str, Any]]:
            return [{"page": 1, "text": "Payment after delivery."}]

    worker = DocumentIndexWorker(
        DocumentIndexJobStore(TEST_DATABASE_URL),
        tenant,
        "test-document-indexer",
        object_store,
        PageIndexAdapter(FakePageIndexClient()),
        lease_seconds=30,
    )
    assert worker.run_once() == "completed"
    assert worker.run_once() == "idle"
    catalog = client.get(f"/api/v1/cases/{case_id}/documents").json()
    assert catalog[0]["index_status"] == "ready"
    with psycopg.connect(TEST_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        assert connection.execute(
            "SELECT status FROM document_index_jobs WHERE tenant_id = %s "
            "AND document_version_id = %s",
            (tenant, uploaded.json()["document_version_id"]),
        ).fetchone() == ("completed",)


def test_document_index_job_fencing_rejects_late_worker_commit(tmp_path: Path) -> None:
    assert TEST_DATABASE_URL is not None and TEST_ADMIN_DATABASE_URL is not None
    tenant, case_id, actor = _case()
    client = _client(tenant, actor, LocalObjectStore(tmp_path))
    uploaded = client.post(
        f"/api/v1/cases/{case_id}/documents",
        headers={"Idempotency-Key": "index-fence-upload-001"},
        data={"role": "primary_contract"},
        files={"file": ("fenced.pdf", _pdf(), PDF_MEDIA_TYPE)},
    )
    assert uploaded.status_code == 201, uploaded.text
    assert _parse_once(tenant, LocalObjectStore(tmp_path)) == "completed"
    store = DocumentIndexJobStore(TEST_DATABASE_URL)
    first = store.claim(tenant, "indexer-before-crash", lease_seconds=30)
    assert first is not None
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        connection.execute(
            "UPDATE document_index_jobs SET lease_until = now() - interval '1 second' "
            "WHERE tenant_id = %s AND id = %s",
            (tenant, first.id),
        )
    second = store.claim(tenant, "indexer-after-crash", lease_seconds=30)
    assert second is not None and second.id == first.id
    assert second.fencing_token == first.fencing_token + 1
    assert not store.complete(tenant, first, "indexer-before-crash", "stale-index-id")
    assert store.complete(tenant, second, "indexer-after-crash", "current-index-id")
    assert client.get(f"/api/v1/cases/{case_id}/documents").json()[0]["index_status"] == "ready"


@pytest.mark.parametrize(
    "use_host_renderer",
    [
        pytest.param(False, id="fake-renderer"),
        pytest.param(
            True,
            id="host-libreoffice",
            marks=pytest.mark.skipif(
                shutil.which("libreoffice") is None, reason="LibreOffice is not installed"
            ),
        ),
    ],
)
def test_docx_indexing_uses_the_persisted_canonical_pdf_render(
    tmp_path: Path,
    use_host_renderer: bool,
) -> None:
    assert TEST_DATABASE_URL is not None
    tenant, case_id, actor = _case()
    object_store = LocalObjectStore(tmp_path)
    client = _client(tenant, actor, object_store)
    docx_data = _docx()
    uploaded = client.post(
        f"/api/v1/cases/{case_id}/documents",
        headers={"Idempotency-Key": "docx-index-upload-0001"},
        data={"role": "primary_contract"},
        files={"file": ("contract.docx", docx_data, DOCX_MEDIA_TYPE)},
    )
    assert uploaded.status_code == 201, uploaded.text
    assert _parse_once(tenant, object_store, use_host_renderer=use_host_renderer) == "completed"
    content_route = (
        f"/api/v1/cases/{case_id}/documents/{uploaded.json()['document_version_id']}/content"
    )
    parsed_artifact = client.get(content_route, params={"kind": "parsed"})
    canonical_render = client.get(content_route, params={"kind": "rendered"})
    assert parsed_artifact.status_code == 200, parsed_artifact.text
    assert canonical_render.status_code == 200, canonical_render.text
    parsed_identity = parsed_artifact.json()
    renderer_identity = parsed_identity["renderer_identity"]
    if use_host_renderer:
        assert renderer_identity.startswith("libreoffice:")
    else:
        assert renderer_identity == "test-renderer:1"
    assert (
        parsed_identity["index_source_sha256"]
        == hashlib.sha256(canonical_render.content).hexdigest()
    )

    class PageIndexClient:
        submitted_pdf: bytes | None = None

        def submit_document(self, file_path: str, *, wait: bool = False) -> dict[str, Any]:
            self.submitted_pdf = Path(file_path).read_bytes()
            assert self.submitted_pdf.startswith(b"%PDF-")
            return {"doc_id": "canonical-docx-pdf"}

        def get_tree(
            self, doc_id: str, node_summary: bool = False, include_text: bool = True
        ) -> dict[str, Any]:
            return {"doc_id": doc_id}

        def get_page_content(self, doc_id: str, pages: str) -> list[dict[str, Any]]:
            return []

    pageindex_client = PageIndexClient()
    worker = DocumentIndexWorker(
        DocumentIndexJobStore(TEST_DATABASE_URL),
        tenant,
        "test-docx-indexer",
        object_store,
        PageIndexAdapter(pageindex_client),
        lease_seconds=30,
    )
    assert worker.run_once() == "completed"
    document = client.get(f"/api/v1/cases/{case_id}/documents").json()[0]
    assert document["index_status"] == "ready"
    assert document["index_error_code"] is None
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        indexed = connection.execute(
            "SELECT rendered_pdf_sha256, renderer_identity FROM document_indexes "
            "WHERE tenant_id = %s AND document_version_id = %s",
            (tenant, UUID(uploaded.json()["document_version_id"])),
        ).fetchone()
    assert indexed is not None
    assert pageindex_client.submitted_pdf is not None
    assert indexed[0] == hashlib.sha256(pageindex_client.submitted_pdf).hexdigest()
    assert indexed[1] == renderer_identity
    selected = client.post(
        f"/api/v1/cases/{case_id}/document-selections",
        json={
            "document_version_id": uploaded.json()["document_version_id"],
            "relation": "main",
            "allow_amount": True,
            "valid_until": (datetime.now(UTC) + timedelta(minutes=30)).isoformat(),
        },
    )
    assert selected.status_code == 201, selected.text
    snapshots = InputSnapshotStore(TEST_DATABASE_URL)
    access = TaskAccess(tenant, case_id, actor, "contract_review", task_requires_amount=True)
    preview = snapshots.preview(access)
    assert len(preview.items) == 1 and preview.items[0].status == "ready"
    assert preview.items[0].rendered_pdf_sha256 == indexed[0]
    assert preview.items[0].renderer_identity == renderer_identity
    frozen = snapshots.freeze(access, preview.id, preview.case_version, preview.manifest_sha256)
    assert UUID(uploaded.json()["document_version_id"]) in frozen.document_version_ids
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        frozen_render = connection.execute(
            "SELECT rendered_pdf_sha256, renderer_identity FROM input_snapshot_documents "
            "WHERE tenant_id = %s AND snapshot_id = %s",
            (tenant, frozen.id),
        ).fetchone()
    assert frozen_render is not None
    assert frozen_render[0] == indexed[0]
    assert frozen_render[1] == indexed[1]


@pytest.mark.skipif(shutil.which("libreoffice") is None, reason="LibreOffice is not installed")
def test_installed_libreoffice_parses_uploaded_docx_in_the_real_worker(tmp_path: Path) -> None:
    assert TEST_DATABASE_URL is not None
    tenant, case_id, actor = _case()
    object_store = LocalObjectStore(tmp_path)
    client = _client(tenant, actor, object_store)
    source = _docx("付款期限为三十日，逾期付款应承担违约责任。")
    uploaded = client.post(
        f"/api/v1/cases/{case_id}/documents",
        headers={"Idempotency-Key": "host-docx-upload-0001"},
        data={"role": "primary_contract"},
        files={"file": ("host-rendered.docx", source, DOCX_MEDIA_TYPE)},
    )
    assert uploaded.status_code == 201, uploaded.text

    worker = DocumentParseWorker(
        DocumentParseJobStore(TEST_DATABASE_URL),
        tenant,
        "host-libreoffice-parser",
        object_store,
        lease_seconds=30,
    )
    assert worker.run_once() == "completed"

    document_id = uploaded.json()["document_version_id"]
    parsed = client.get(f"/api/v1/cases/{case_id}/documents/{document_id}/content?kind=parsed")
    rendered = client.get(f"/api/v1/cases/{case_id}/documents/{document_id}/content?kind=rendered")
    assert parsed.status_code == 200, parsed.text
    assert rendered.status_code == 200 and rendered.content.startswith(b"%PDF-")
    parsed_artifact = parsed.json()
    assert parsed_artifact["source_sha256"] == hashlib.sha256(source).hexdigest()
    assert parsed_artifact["index_source_sha256"] == hashlib.sha256(rendered.content).hexdigest()
    assert parsed_artifact["renderer_identity"].startswith("libreoffice:")
    assert "付款期限为三十日" in "\n".join(parsed_artifact["pages"])


def test_upload_rolls_back_objects_when_database_registration_fails(tmp_path: Path) -> None:
    assert TEST_DATABASE_URL is not None and TEST_ADMIN_DATABASE_URL is not None
    tenant, case_id, actor = _case()
    object_store = LocalObjectStore(tmp_path / "objects")
    trigger_name = f"fail_index_insert_{uuid4().hex}"
    function_name = f"{trigger_name}_fn"
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute(
            f"CREATE FUNCTION {function_name}() RETURNS trigger LANGUAGE plpgsql AS "
            f"$$ BEGIN IF NEW.case_id = '{case_id}'::uuid THEN "
            "RAISE EXCEPTION 'test-only-index-registration-failure'; END IF; "
            "RETURN NEW; END $$"
        )
        connection.execute(
            f"CREATE TRIGGER {trigger_name} BEFORE INSERT ON document_parse_jobs "
            f"FOR EACH ROW EXECUTE FUNCTION {function_name}()"
        )
    try:
        ingestion = DocumentIngestionStore(TEST_DATABASE_URL, object_store)
        with pytest.raises(psycopg.Error):
            ingestion.upload(
                tenant,
                case_id,
                actor,
                "rollback.pdf",
                PDF_MEDIA_TYPE,
                _pdf("rollback compensation"),
                "primary_contract",
                "rollback-upload-0001",
            )
    finally:
        with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
            connection.execute(f"DROP TRIGGER {trigger_name} ON document_parse_jobs")
            connection.execute(f"DROP FUNCTION {function_name}()")

    assert _uploaded_count(tenant, case_id) == 0
    assert not [path for path in (tmp_path / "objects").rglob("*") if path.is_file()]


def test_object_gc_deletes_only_old_unreferenced_objects_for_configured_tenant(
    tmp_path: Path,
) -> None:
    assert TEST_DATABASE_URL is not None
    tenant, case_id, actor = _case()
    object_store = LocalObjectStore(tmp_path / "objects")
    client = _client(tenant, actor, object_store)
    source = _pdf("retained referenced source")
    uploaded = client.post(
        f"/api/v1/cases/{case_id}/documents",
        headers={"Idempotency-Key": "gc-referenced-upload-01"},
        data={"role": "primary_contract"},
        files={"file": ("retained.pdf", source, PDF_MEDIA_TYPE)},
    )
    assert uploaded.status_code == 201, uploaded.text
    orphan = object_store.put(tenant, case_id, b"unreferenced orphan bytes")
    other_tenant_orphan = object_store.put(uuid4(), case_id, b"other tenant orphan")
    temporary_digest = hashlib.sha256(b"interrupted upload").hexdigest()
    stale_temporary = (
        object_store.root / str(tenant) / str(case_id) / f".upload-{temporary_digest}-interrupted"
    )
    stale_temporary.write_bytes(b"partial upload bytes")
    other_tenant_temp = (
        object_store.root
        / str(other_tenant_orphan.tenant_id)
        / str(case_id)
        / f".upload-{temporary_digest}-other-tenant"
    )
    other_tenant_temp.write_bytes(b"other tenant partial bytes")

    candidates = list(object_store.iter_before(tenant, datetime.now(UTC) + timedelta(days=1)))
    old_time = (datetime.now(UTC) - timedelta(days=2)).timestamp()
    for candidate in candidates:
        path = (
            object_store.root
            / str(candidate.reference.tenant_id)
            / str(candidate.reference.case_id)
            / candidate.reference.sha256
        )
        os.utime(path, (old_time, old_time))
    os.utime(stale_temporary, (old_time, old_time))
    os.utime(other_tenant_temp, (old_time, old_time))

    result = DocumentObjectGarbageCollector(
        TEST_DATABASE_URL,
        tenant,
        object_store,
        minimum_age=timedelta(days=1),
    ).sweep_once()

    assert (result.scanned, result.deleted, result.retained, result.already_missing) == (3, 2, 1, 0)
    assert not stale_temporary.exists()
    assert other_tenant_temp.exists()
    with pytest.raises(FileNotFoundError):
        object_store.get(orphan)
    assert object_store.get(other_tenant_orphan) == b"other tenant orphan"
    assert (
        object_store.get(ObjectRef(tenant, case_id, uploaded.json()["content_hash"], len(source)))
        == source
    )


def test_parse_job_lease_expiry_is_fenced_and_reclaimed(tmp_path: Path) -> None:
    assert TEST_DATABASE_URL is not None and TEST_ADMIN_DATABASE_URL is not None
    tenant, case_id, actor = _case()
    object_store = LocalObjectStore(tmp_path)
    client = _client(tenant, actor, object_store)
    uploaded = client.post(
        f"/api/v1/cases/{case_id}/documents",
        headers={"Idempotency-Key": "parse-fence-upload-0001"},
        data={"role": "primary_contract"},
        files={"file": ("parse-fence.pdf", _pdf(), PDF_MEDIA_TYPE)},
    )
    assert uploaded.status_code == 201
    store = DocumentParseJobStore(TEST_DATABASE_URL)
    first = store.claim(tenant, "parser-before-crash", lease_seconds=30)
    assert first is not None
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        connection.execute(
            "UPDATE document_parse_jobs SET lease_until = now() - interval '1 second' "
            "WHERE tenant_id = %s AND id = %s",
            (tenant, first.id),
        )
    second = store.claim(tenant, "parser-after-crash", lease_seconds=30)
    assert second is not None and second.id == first.id
    assert second.fencing_token == first.fencing_token + 1
    assert not store.heartbeat(tenant, first, "parser-before-crash", lease_seconds=30)
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        connection.execute(
            "UPDATE document_parse_jobs SET lease_until = now() - interval '1 second' "
            "WHERE tenant_id = %s AND id = %s",
            (tenant, second.id),
        )
    assert _parse_once(tenant, object_store) == "completed"
    with psycopg.connect(TEST_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        state = connection.execute(
            "SELECT j.status, j.fencing_token, d.parse_status FROM document_parse_jobs j "
            "JOIN document_versions d ON d.tenant_id = j.tenant_id "
            "AND d.id = j.document_version_id WHERE j.tenant_id = %s AND j.id = %s",
            (tenant, first.id),
        ).fetchone()
    assert state == ("completed", second.fencing_token + 1, "completed")


def test_upload_denial_and_parse_failure_are_isolated_and_visible(
    tmp_path: Path,
) -> None:
    assert TEST_DATABASE_URL is not None
    tenant, case_id, actor = _case()
    store = LocalObjectStore(tmp_path)
    app = create_app()
    app.dependency_overrides[require_principal] = lambda: Principal(tenant, uuid4())
    app.dependency_overrides[get_document_ingestion_store] = lambda: DocumentIngestionStore(
        TEST_DATABASE_URL, store
    )
    app.dependency_overrides[get_document_catalog_store] = lambda: DocumentCatalogStore(
        TEST_DATABASE_URL
    )
    client = TestClient(app)
    route = f"/api/v1/cases/{case_id}/documents"
    denied = client.post(
        route,
        headers={"Idempotency-Key": "upload-00000003"},
        data={"role": "primary_contract"},
        files={"file": ("contract.pdf", _pdf(), PDF_MEDIA_TYPE)},
    )
    assert denied.status_code == 403
    assert not (tmp_path / str(tenant)).exists()

    app.dependency_overrides[require_principal] = lambda: Principal(tenant, actor)
    invalid = client.post(
        route,
        headers={"Idempotency-Key": "upload-00000004"},
        data={"role": "primary_contract"},
        files={"file": ("broken.pdf", b"not a pdf", PDF_MEDIA_TYPE)},
    )
    assert invalid.status_code == 201
    assert invalid.json()["parse_status"] == "queued"
    assert _uploaded_count(tenant, case_id) == 1
    assert _parse_once(tenant, store) == "failed"
    failed = client.get(route).json()[0]
    assert failed["parse_status"] == "failed"
    assert failed["parse_error_code"] == "pdf_signature_invalid"
    assert failed["warnings"] == []


def test_docx_render_retry_exposes_only_a_safe_parse_error_code(tmp_path: Path) -> None:
    assert TEST_DATABASE_URL is not None and TEST_ADMIN_DATABASE_URL is not None
    tenant, case_id, actor = _case()
    object_store = LocalObjectStore(tmp_path)
    client = _client(tenant, actor, object_store)
    uploaded = client.post(
        f"/api/v1/cases/{case_id}/documents",
        headers={"Idempotency-Key": "docx-render-retry-0001"},
        data={"role": "primary_contract"},
        files={"file": ("retry.docx", _docx(), DOCX_MEDIA_TYPE)},
    )
    assert uploaded.status_code == 201, uploaded.text

    def timed_out_renderer(_data: bytes) -> RenderedDocx:
        raise DocxRenderError("docx_renderer_timeout")

    worker = DocumentParseWorker(
        DocumentParseJobStore(TEST_DATABASE_URL),
        tenant,
        "test-docx-retry-parser",
        object_store,
        docx_renderer=timed_out_renderer,
        lease_seconds=30,
        retry_delay_seconds=0,
    )
    assert worker.run_once() == "queued"
    document_id = uploaded.json()["document_version_id"]
    listed = client.get(f"/api/v1/cases/{case_id}/documents").json()[0]
    assert listed["document_version_id"] == document_id
    assert listed["parse_status"] == "queued"
    assert listed["parse_error_code"] == "docx_renderer_timeout"

    with psycopg.connect(TEST_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        retry_state = connection.execute(
            "SELECT job.status, job.last_error, document.parse_status "
            "FROM document_parse_jobs AS job JOIN document_versions AS document "
            "ON document.tenant_id = job.tenant_id "
            "AND document.id = job.document_version_id "
            "WHERE job.tenant_id = %s AND job.document_version_id = %s",
            (tenant, UUID(document_id)),
        ).fetchone()
    assert retry_state == ("queued", "docx_renderer_timeout", "queued")

    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        connection.execute(
            "UPDATE document_parse_jobs SET last_error = %s "
            "WHERE tenant_id = %s AND document_version_id = %s",
            ("/private/path/secret.docx: raw converter diagnostic", tenant, UUID(document_id)),
        )
        connection.execute(
            "UPDATE document_versions SET warnings = %s::jsonb WHERE tenant_id = %s AND id = %s",
            (
                json.dumps(["/private/path/secret.docx: raw converter diagnostic"]),
                tenant,
                UUID(document_id),
            ),
        )
    listed_after_untrusted_error = client.get(f"/api/v1/cases/{case_id}/documents").json()[0]
    assert listed_after_untrusted_error["parse_error_code"] is None
    assert listed_after_untrusted_error["warnings"] == []


def test_upload_route_requires_trusted_identity(tmp_path: Path) -> None:
    tenant, case_id, actor = _case()
    app = create_app()
    app.dependency_overrides[get_document_ingestion_store] = lambda: DocumentIngestionStore(
        TEST_DATABASE_URL or "", LocalObjectStore(tmp_path)
    )
    client = TestClient(app)
    response = client.post(
        f"/api/v1/cases/{case_id}/documents",
        headers={"Idempotency-Key": "upload-00000005"},
        data={"role": "primary_contract", "tenant_id": str(uuid4())},
        files={"file": ("contract.pdf", _pdf(), PDF_MEDIA_TYPE)},
    )
    assert response.status_code == 401
    assert response.headers["content-type"] == "application/problem+json"
    assert "tenant_id" not in response.text


def test_primary_upload_precedes_entity_confirmation_but_scope_waits(tmp_path: Path) -> None:
    assert TEST_DATABASE_URL is not None and TEST_ADMIN_DATABASE_URL is not None
    tenant, actor, entity = uuid4(), uuid4(), uuid4()
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        case_id = connection.execute(
            "INSERT INTO audit_cases (tenant_id, title) VALUES (%s, '主体待确认案件') RETURNING id",
            (tenant,),
        ).fetchone()[0]
        connection.execute(
            "INSERT INTO case_members (tenant_id, case_id, actor_id, role) "
            "VALUES (%s, %s, %s, 'reviewer')",
            (tenant, case_id, actor),
        )
    store = LocalObjectStore(tmp_path)
    client = _client(tenant, actor, store)
    route = f"/api/v1/cases/{case_id}/documents"
    primary = client.post(
        route,
        headers={"Idempotency-Key": "unconfirmed-main-0001"},
        data={"role": "primary_contract"},
        files={"file": ("main.pdf", _pdf(), PDF_MEDIA_TYPE)},
    )
    assert primary.status_code == 201, primary.text
    assert primary.json()["parse_status"] == "queued"
    attachment = client.post(
        route,
        headers={"Idempotency-Key": "unconfirmed-attachment-01"},
        data={"role": "attachment"},
        files={"file": ("appendix.pdf", _pdf("appendix"), PDF_MEDIA_TYPE)},
    )
    assert attachment.status_code == 409
    selected = client.post(
        f"/api/v1/cases/{case_id}/document-selections",
        json={
            "document_version_id": primary.json()["document_version_id"],
            "relation": "main",
            "allow_amount": False,
            "valid_until": (datetime.now(UTC) + timedelta(minutes=20)).isoformat(),
        },
    )
    assert selected.status_code == 409
    with pytest.raises(PreviewConflict, match="case_legal_entity_unconfirmed"):
        InputSnapshotStore(TEST_DATABASE_URL).preview(
            TaskAccess(tenant, case_id, actor, "contract_review")
        )

    confirmation_store = CaseLegalEntityStore(TEST_DATABASE_URL)
    confirmed = confirmation_store.confirm(
        tenant, case_id, actor, entity, 2, "已按签约主体资料核对", "entity-confirm-00001"
    )
    assert confirmed.legal_entity_id == entity and confirmed.version == 3
    with psycopg.connect(TEST_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        document_entity = connection.execute(
            "SELECT legal_entity_id FROM document_versions WHERE tenant_id = %s AND id = %s",
            (tenant, UUID(primary.json()["document_version_id"])),
        ).fetchone()[0]
    assert document_entity == entity
    attachment_after = client.post(
        route,
        headers={"Idempotency-Key": "confirmed-attachment-001"},
        data={"role": "attachment"},
        files={"file": ("appendix.pdf", _pdf("appendix"), PDF_MEDIA_TYPE)},
    )
    assert attachment_after.status_code == 201, attachment_after.text


def test_reviewer_selects_main_and_attachment_with_explicit_amount_scope(
    tmp_path: Path,
) -> None:
    assert TEST_DATABASE_URL is not None
    assert TEST_ADMIN_DATABASE_URL is not None
    tenant, case_id, actor = _case()
    client = _client(tenant, actor, LocalObjectStore(tmp_path))
    upload_route = f"/api/v1/cases/{case_id}/documents"
    uploaded: dict[str, dict[str, Any]] = {}
    for role, name in (("primary_contract", "main.pdf"), ("attachment", "annex.pdf")):
        response = client.post(
            upload_route,
            headers={"Idempotency-Key": f"select-{role}-0001"},
            data={"role": role},
            files={"file": (name, _pdf(name), PDF_MEDIA_TYPE)},
        )
        assert response.status_code == 201, response.text
        uploaded[role] = response.json()
    assert _parse_once(tenant, LocalObjectStore(tmp_path)) == "completed"
    assert _parse_once(tenant, LocalObjectStore(tmp_path)) == "completed"

    select_route = f"/api/v1/cases/{case_id}/document-selections"
    valid_until = datetime.now(UTC) + timedelta(minutes=30)
    main_request = {
        "document_version_id": uploaded["primary_contract"]["document_version_id"],
        "relation": "main",
        "allow_amount": True,
        "valid_until": valid_until.isoformat(),
    }
    main = client.post(select_route, json=main_request)
    assert main.status_code == 201, main.text
    replay = client.post(select_route, json=main_request)
    assert replay.status_code == 201
    assert replay.json() == main.json()
    assert main.json()["purpose"] == "contract_review"
    assert main.json()["allow_text"] is True
    assert main.json()["allow_amount"] is True
    assert main.json()["case_version"] == 4

    attachment = client.post(
        select_route,
        json={
            **main_request,
            "document_version_id": uploaded["attachment"]["document_version_id"],
            "relation": "attachment",
            "allow_amount": True,
        },
    )
    assert attachment.status_code == 201, attachment.text
    assert attachment.json()["case_version"] == 5

    with psycopg.connect(TEST_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        assert (
            connection.execute(
                "SELECT count(*) FROM contract_selections WHERE tenant_id = %s AND case_id = %s",
                (tenant, case_id),
            ).fetchone()[0]
            == 2
        )
        assert (
            connection.execute(
                "SELECT count(*) FROM document_grants WHERE tenant_id = %s "
                "AND purpose = 'contract_review' AND allow_text",
                (tenant,),
            ).fetchone()[0]
            == 2
        )
        assert (
            connection.execute(
                "SELECT count(*) FROM document_grants WHERE tenant_id = %s "
                "AND purpose = 'contract_review' AND NOT allow_amount",
                (tenant,),
            ).fetchone()[0]
            == 0
        )
        assert (
            connection.execute(
                "SELECT count(*) FROM document_versions WHERE tenant_id = %s "
                "AND uploaded_case_id = %s AND status = 'active'",
                (tenant, case_id),
            ).fetchone()[0]
            == 2
        )

    preview = InputSnapshotStore(TEST_DATABASE_URL).preview(
        TaskAccess(tenant, case_id, actor, "contract_review", task_requires_amount=True)
    )
    assert len(preview.items) == 2
    assert {item.relation for item in preview.items} == {"main", "attachment"}
    assert {item.relation: item.status for item in preview.items} == {
        "main": "index_unavailable",
        "attachment": "index_unavailable",
    }

    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        documents = connection.execute(
            "SELECT id, source_sha256 FROM document_versions "
            "WHERE tenant_id = %s AND uploaded_case_id = %s",
            (tenant, case_id),
        ).fetchall()
        for document_id, source_hash in documents:
            connection.execute(
                "INSERT INTO document_indexes "
                "(tenant_id, document_version_id, source_sha256, pageindex_doc_id) "
                "VALUES (%s, %s, %s, %s)",
                (tenant, document_id, source_hash, f"fake-{document_id}"),
            )
    ready_preview = InputSnapshotStore(TEST_DATABASE_URL).preview(
        TaskAccess(tenant, case_id, actor, "contract_review", task_requires_amount=True)
    )
    assert {item.status for item in ready_preview.items} == {"ready"}
    frozen = InputSnapshotStore(TEST_DATABASE_URL).freeze(
        TaskAccess(tenant, case_id, actor, "contract_review", task_requires_amount=True),
        ready_preview.id,
        ready_preview.case_version,
        ready_preview.manifest_sha256,
    )
    assert frozen.document_version_ids == {
        UUID(uploaded["primary_contract"]["document_version_id"]),
        UUID(uploaded["attachment"]["document_version_id"]),
    }


def test_only_reviewer_can_authorize_and_existing_grant_cannot_be_broadened(
    tmp_path: Path,
) -> None:
    assert TEST_ADMIN_DATABASE_URL is not None
    tenant, case_id, actor = _case()
    client = _client(tenant, actor, LocalObjectStore(tmp_path))
    file_content = _pdf()
    upload = client.post(
        f"/api/v1/cases/{case_id}/documents",
        headers={"Idempotency-Key": "select-auth-00001"},
        data={"role": "primary_contract"},
        files={"file": ("main.pdf", file_content, PDF_MEDIA_TYPE)},
    )
    assert _parse_once(tenant, LocalObjectStore(tmp_path)) == "completed"
    document_id = upload.json()["document_version_id"]
    route = f"/api/v1/cases/{case_id}/document-selections"
    request = {
        "document_version_id": document_id,
        "relation": "main",
        "allow_amount": False,
        "valid_until": (datetime.now(UTC) + timedelta(minutes=20)).isoformat(),
    }
    submitter = uuid4()
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        connection.execute(
            "INSERT INTO case_members (tenant_id, case_id, actor_id, role) "
            "VALUES (%s, %s, %s, 'submitter')",
            (tenant, case_id, submitter),
        )
    client.app.dependency_overrides[require_principal] = lambda: Principal(tenant, submitter)
    submitter_denied = client.post(route, json=request)
    assert submitter_denied.status_code == 403

    client.app.dependency_overrides[require_principal] = lambda: Principal(tenant, uuid4())
    denied = client.post(route, json=request)
    assert denied.status_code == 403
    assert "reviewer_membership_required" not in denied.text

    client.app.dependency_overrides[require_principal] = lambda: Principal(tenant, actor)
    selected = client.post(route, json=request)
    assert selected.status_code == 201, selected.text
    upload_replay = client.post(
        f"/api/v1/cases/{case_id}/documents",
        headers={"Idempotency-Key": "select-auth-00001"},
        data={"role": "primary_contract"},
        files={"file": ("main.pdf", file_content, PDF_MEDIA_TYPE)},
    )
    assert upload_replay.status_code == 201
    assert upload_replay.json()["document_status"] == "active"
    assert upload_replay.json()["authorization_status"] == "not_granted_by_upload"
    broadened = client.post(route, json={**request, "allow_amount": True})
    assert broadened.status_code == 409
    assert "selection_already_authorized_differently" not in broadened.text
    wrong_relation = client.post(route, json={**request, "relation": "attachment"})
    assert wrong_relation.status_code == 409


def test_revocation_is_audited_idempotent_and_invalidates_bound_runs(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.delenv("CONTRACT_REVIEW_EXECUTION_ENABLED", raising=False)
    assert TEST_DATABASE_URL is not None
    assert TEST_ADMIN_DATABASE_URL is not None
    tenant, case_id, actor = _case()
    object_store = LocalObjectStore(tmp_path)
    client = _client(tenant, actor, object_store)
    source_data = _pdf("Payment amount USD 1000, due in 30 days.")
    uploaded = client.post(
        f"/api/v1/cases/{case_id}/documents",
        headers={"Idempotency-Key": "revoke-upload-0001"},
        data={"role": "primary_contract"},
        files={"file": ("main.pdf", source_data, PDF_MEDIA_TYPE)},
    ).json()
    assert _parse_once(tenant, object_store) == "completed"
    selection_route = f"/api/v1/cases/{case_id}/document-selections"
    valid_until = datetime.now(UTC) + timedelta(minutes=45)
    selection_request = {
        "document_version_id": uploaded["document_version_id"],
        "relation": "main",
        "allow_amount": True,
        "valid_until": valid_until.isoformat(),
    }
    selected = client.post(selection_route, json=selection_request)
    assert selected.status_code == 201, selected.text
    selection_id = selected.json()["selection_id"]

    second_reviewer = uuid4()
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        connection.execute(
            "INSERT INTO case_members (tenant_id, case_id, actor_id, role) "
            "VALUES (%s, %s, %s, 'reviewer')",
            (tenant, case_id, second_reviewer),
        )
        source = connection.execute(
            "SELECT id, source_sha256 FROM document_versions "
            "WHERE tenant_id = %s AND uploaded_case_id = %s",
            (tenant, case_id),
        ).fetchone()
        assert source is not None
        connection.execute(
            "INSERT INTO document_indexes "
            "(tenant_id, document_version_id, source_sha256, pageindex_doc_id) "
            "VALUES (%s, %s, %s, %s)",
            (tenant, source[0], source[1], f"revoke-index-{source[0]}"),
        )
    client.app.dependency_overrides[require_principal] = lambda: Principal(tenant, second_reviewer)
    second_selection = client.post(
        selection_route,
        json={**selection_request, "allow_amount": False},
    )
    assert second_selection.status_code == 201, second_selection.text
    assert second_selection.json()["selection_id"] == selection_id

    access = TaskAccess(tenant, case_id, actor, "contract_review", task_requires_amount=True)
    snapshot_store = InputSnapshotStore(TEST_DATABASE_URL)
    preview = snapshot_store.preview(access)
    assert preview.items[0].status == "ready"
    frozen = snapshot_store.freeze(
        access, preview.id, preview.case_version, preview.manifest_sha256
    )
    bound_access = TaskAccess(
        tenant,
        case_id,
        actor,
        "contract_review",
        task_requires_amount=True,
        input_snapshot_id=frozen.id,
    )
    jobs = JobStore(TEST_DATABASE_URL)
    client.app.dependency_overrides[require_principal] = lambda: Principal(tenant, actor)
    start_route = f"/api/v1/cases/{case_id}/runs"
    start_headers = {"Idempotency-Key": "revoke-bound-run-001"}
    start_body = {"input_snapshot_id": str(frozen.id)}
    disabled = client.post(start_route, json=start_body, headers=start_headers)
    assert disabled.status_code == 503
    assert disabled.json()["code"] == "review_execution_disabled"
    monkeypatch.setenv("APP_ENV", "production")
    monkeypatch.setenv("CONTRACT_REVIEW_EXECUTION_ENABLED", "true")
    still_disabled = client.post(start_route, json=start_body, headers=start_headers)
    assert still_disabled.status_code == 503
    assert still_disabled.json()["code"] == "review_execution_disabled"
    monkeypatch.setenv("APP_ENV", "development")
    monkeypatch.setenv("CONTRACT_REVIEW_MODEL", "offline-mock")
    development_started = client.post(start_route, json=start_body, headers=start_headers)
    assert development_started.status_code == 202, development_started.text
    monkeypatch.setenv("APP_ENV", "test")
    started = client.post(start_route, json=start_body, headers=start_headers)
    assert started.status_code == 202, started.text
    assert started.json() == development_started.json()
    replayed = client.post(start_route, json=start_body, headers=start_headers)
    assert replayed.status_code == 202 and replayed.json() == started.json()
    assert client.get(f"/api/v1/runs/{started.json()['run_id']}").json() == started.json()
    assert client.get(f"/api/v1/runs/{started.json()['run_id']}/results").json()["results"] == []
    run_page = client.get(f"{start_route}?limit=1")
    assert run_page.status_code == 200
    assert run_page.json()["items"] == [started.json()]
    assert run_page.json()["next_cursor"] is None
    client.app.dependency_overrides[require_principal] = lambda: Principal(tenant, uuid4())
    assert client.get(start_route).status_code == 403
    client.app.dependency_overrides[require_principal] = lambda: Principal(tenant, actor)
    assert (
        client.post(
            start_route,
            json={**start_body, "document_version_id": str(uploaded["document_version_id"])},
            headers={"Idempotency-Key": "forged-run-000001"},
        ).status_code
        == 422
    )
    run_id = UUID(started.json()["run_id"])
    with psycopg.connect(TEST_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        queued = connection.execute(
            "SELECT task_type, idempotency_key, payload FROM tasks "
            "WHERE tenant_id = %s AND run_id = %s",
            (tenant, run_id),
        ).fetchone()
        assert queued == (
            "contract_review",
            start_headers["Idempotency-Key"],
            {"input_snapshot_id": str(frozen.id)},
        )
    leased_task = jobs.claim(tenant, "review-worker", lease_seconds=30)
    assert leased_task is not None and leased_task.run_id == run_id
    task_id = leased_task.id
    amount_candidate_evidence_id = uuid4()
    result_payload = {
        "results": [
            {
                "result_id": str(uuid4()),
                "rule_id": "legal.payment_conditions",
                "rule_version": "1.0.0",
                "judgment": "undetermined",
                "severity": "undetermined",
                "summary": "Needs further review.",
                "evidence": [
                    {
                        "evidence_id": str(amount_candidate_evidence_id),
                        "kind": "document_span",
                        "support": "direct",
                        "source_version": str(uploaded["document_version_id"]),
                        "document_version_id": uploaded["document_version_id"],
                        "page": 1,
                        "start_offset": 0,
                        "end_offset": 7,
                        "quote": "Payment",
                        "text_hash": "a" * 64,
                    }
                ],
                "slot_candidates": [
                    {
                        "slot_id": "payment_schedule",
                        "state": "filled",
                        "value": [
                            {
                                "amount_type": "fixed",
                                "amount": "1000",
                                "currency": "USD",
                            }
                        ],
                        "fact_type": "contract.payment_schedule",
                        "evidence_ids": [str(amount_candidate_evidence_id)],
                    }
                ],
                "limitations": ["Synthetic integration-test result."],
                "created_at": datetime.now(UTC).isoformat(),
            },
            {
                "result_id": str(uuid4()),
                "rule_id": "legal.payment_conditions",
                "rule_version": "1.0.0",
                "judgment": "undetermined",
                "severity": "undetermined",
                "summary": "No validated candidate found.",
                "evidence": [],
                "limitations": ["Absence of a candidate is not proof of absence."],
                "created_at": datetime.now(UTC).isoformat(),
            },
        ]
    }
    serialized_result = json.dumps(
        result_payload, sort_keys=True, ensure_ascii=False, separators=(",", ":")
    )
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        connection.execute(
            "INSERT INTO task_results (tenant_id, task_id, run_id, result, content_hash) "
            "VALUES (%s, %s, %s, %s::jsonb, %s)",
            (
                tenant,
                task_id,
                run_id,
                serialized_result,
                hashlib.sha256(serialized_result.encode()).hexdigest(),
            ),
        )
    result_response = client.get(f"/api/v1/runs/{run_id}/results")
    assert result_response.status_code == 200, result_response.text
    assert len(result_response.json()["results"]) == 2
    assert (
        result_response.json()["results"][0]["slot_candidates"][0]["value"][0]["amount"] == "1000"
    )
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        revoked_amount = connection.execute(
            "UPDATE document_grants SET allow_amount = false "
            "WHERE tenant_id = %s AND actor_id = %s AND document_version_id = %s "
            "AND purpose = 'contract_review'",
            (tenant, actor, uploaded["document_version_id"]),
        )
        assert revoked_amount.rowcount == 1
    assert client.get(f"/api/v1/runs/{run_id}/results").status_code == 403
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        restored_amount = connection.execute(
            "UPDATE document_grants SET allow_amount = true "
            "WHERE tenant_id = %s AND actor_id = %s AND document_version_id = %s "
            "AND purpose = 'contract_review'",
            (tenant, actor, uploaded["document_version_id"]),
        )
        assert restored_amount.rowcount == 1
    assert client.get(f"/api/v1/runs/{run_id}/results").status_code == 200
    client.app.dependency_overrides[require_principal] = lambda: Principal(tenant, second_reviewer)
    assert client.get(f"/api/v1/runs/{run_id}/results").status_code == 403
    client.app.dependency_overrides[require_principal] = lambda: Principal(tenant, actor)
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        connection.execute(
            "DELETE FROM task_results WHERE tenant_id = %s AND task_id = %s",
            (tenant, task_id),
        )

    revoke_route = f"/api/v1/cases/{case_id}/document-selections/{selection_id}/revoke"
    revoke_request = {
        "expected_case_version": frozen.case_version,
        "reason": "Review scope changed",
    }
    headers = {"Idempotency-Key": "revoke-command-0001"}
    revoked = client.post(revoke_route, json=revoke_request, headers=headers)
    assert revoked.status_code == 200, revoked.text
    assert revoked.json()["revoked_grant_count"] == 2
    assert revoked.json()["invalidated_run_count"] == 1
    assert revoked.json()["case_version"] == frozen.case_version + 1
    assert client.post(revoke_route, json=revoke_request, headers=headers).json() == revoked.json()
    access_events_route = f"/api/v1/cases/{case_id}/document-access-events"
    first_access_page = client.get(access_events_route, params={"limit": 1})
    assert first_access_page.status_code == 200, first_access_page.text
    assert len(first_access_page.json()["items"]) == 1
    assert first_access_page.json()["items"][0]["action"] == "selection_revoked"
    assert first_access_page.json()["items"][0]["reason"] == "Review scope changed"
    assert first_access_page.json()["items"][0]["file_name"] == "main.pdf"
    next_access_page = client.get(
        access_events_route,
        params={"limit": 1, "cursor": first_access_page.json()["next_cursor"]},
    )
    assert next_access_page.status_code == 200, next_access_page.text
    assert next_access_page.json()["items"][0]["action"] == "selection_granted"
    client.app.dependency_overrides[require_principal] = lambda: Principal(tenant, uuid4())
    assert client.get(access_events_route).status_code == 403
    client.app.dependency_overrides[require_principal] = lambda: Principal(tenant, actor)
    catalog_state = client.get(f"/api/v1/cases/{case_id}/documents").json()[0]
    assert catalog_state["selection_id"] == str(selection_id)
    assert catalog_state["selection_status"] == "revoked"
    assert catalog_state["authorization_status"] == "not_authorized"
    assert catalog_state["authorization_valid_until"] is None
    assert not jobs.complete(tenant, leased_task, "review-worker", {"status": "passed"})
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        connection.execute(
            "INSERT INTO task_results (tenant_id, task_id, run_id, result, content_hash) "
            "VALUES (%s, %s, %s, %s::jsonb, %s)",
            (
                tenant,
                task_id,
                run_id,
                serialized_result,
                hashlib.sha256(serialized_result.encode()).hexdigest(),
            ),
        )
    assert client.get(f"/api/v1/runs/{run_id}/results").status_code == 403
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        connection.execute(
            "DELETE FROM task_results WHERE tenant_id = %s AND task_id = %s",
            (tenant, task_id),
        )

    with psycopg.connect(TEST_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        assert (
            connection.execute(
                "SELECT status FROM audit_runs WHERE tenant_id = %s AND id = %s",
                (tenant, run_id),
            ).fetchone()[0]
            == "cancelled"
        )
        assert connection.execute(
            "SELECT status, fencing_token FROM tasks WHERE tenant_id = %s AND id = %s",
            (tenant, task_id),
        ).fetchone() == ("cancelled", 2)
        assert (
            connection.execute(
                "SELECT count(*) FROM task_results WHERE tenant_id = %s AND task_id = %s",
                (tenant, task_id),
            ).fetchone()[0]
            == 0
        )
        audit = connection.execute(
            "SELECT action, reason, case_version, permissions, invalidated_run_ids "
            "FROM document_access_events WHERE tenant_id = %s AND idempotency_key = %s",
            (tenant, headers["Idempotency-Key"]),
        ).fetchone()
        assert audit is not None
        assert audit[0:3] == (
            "selection_revoked",
            "Review scope changed",
            revoked.json()["case_version"],
        )
        assert len(audit[3]) == 2
        assert audit[4] == [run_id]

    projector = AuthorizationProjector(TEST_DATABASE_URL)
    with pytest.raises(RetrievalDenied, match="result_source_not_authorized"):
        projector.require_result_sources(
            bound_access, frozenset({UUID(uploaded["document_version_id"])})
        )
    with pytest.raises(BoundRunConflict, match="input_snapshot_stale_or_scope_changed"):
        jobs.create_bound_run(bound_access, "revoke-bound-run-002")


@pytest.mark.parametrize(
    "price_case", ("normal", "tax_conflict", "amount_conflict", "missing_denominator")
)
def test_contract_review_worker_emits_only_validated_undetermined_candidates(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    price_case: str,
) -> None:
    monkeypatch.setenv("APP_ENV", "test")
    monkeypatch.setenv("CONTRACT_REVIEW_EXECUTION_ENABLED", "true")
    monkeypatch.setenv("CONTRACT_REVIEW_MODEL", "offline-mock")
    pytest.importorskip("agentscope")
    from agentscope.credential import OpenAICredential
    from agentscope.formatter import OpenAIChatFormatter
    from agentscope.message import TextBlock, ToolCallBlock
    from agentscope.model import ChatModelBase, ChatResponse, ChatUsage
    from pydantic import SecretStr

    from contract_app.workflow.contract_review import ContractReviewProcessor
    from contract_app.workflow.worker import Worker

    assert TEST_DATABASE_URL is not None
    tenant, case_id, actor = _case()
    object_store = LocalObjectStore(tmp_path)
    client = _client(tenant, actor, object_store)
    # The tiny PDF fixture uses built-in Helvetica; keep its text ASCII so the
    # parser can extract it deterministically without bundling a CJK font.
    quote = (
        "Service PRC customer CNY 1000 tax 100 net 900; attachment CNY 1200; pays 100% acceptance."
        if price_case == "amount_conflict"
        else "Service PRC customer CNY 1000 tax 100 net 900; pays 100% acceptance."
    )
    tax_amount = "90" if price_case == "tax_conflict" else "100"
    price_source_candidates = [
        {
            "source_type": "contract_body",
            "amount": "1000",
            "currency": "CNY",
            "priority": "primary",
        }
    ]
    if price_case == "amount_conflict":
        price_source_candidates.append(
            {
                "source_type": "attachment",
                "amount": "1200",
                "currency": "CNY",
                "priority": "secondary",
            }
        )
    payment_denominator = "unspecified" if price_case == "missing_denominator" else "contract_total"
    fixed_payment_amount = "900" if price_case == "amount_conflict" else "1000"
    upload = client.post(
        f"/api/v1/cases/{case_id}/documents",
        headers={"Idempotency-Key": "review-worker-upload-001"},
        data={"role": "primary_contract"},
        files={"file": ("payment.pdf", _pdf(quote), PDF_MEDIA_TYPE)},
    )
    assert upload.status_code == 201, upload.text
    assert _parse_once(tenant, object_store) == "completed"
    document_id = upload.json()["document_version_id"]
    price_source_evidence = [
        {
            "document_version_id": document_id,
            "page": 1,
            "quote": "Service PRC customer CNY 1000 tax 100 net 900;",
        }
    ]
    if price_case == "amount_conflict":
        price_source_evidence.append(
            {
                "document_version_id": document_id,
                "page": 1,
                "quote": "attachment CNY 1200",
            }
        )
    selected = client.post(
        f"/api/v1/cases/{case_id}/document-selections",
        json={
            "document_version_id": document_id,
            "relation": "main",
            "allow_amount": True,
            "valid_until": (datetime.now(UTC) + timedelta(hours=1)).isoformat(),
        },
    )
    assert selected.status_code == 201, selected.text

    class SyntheticPageIndex:
        def submit_document(self, file_path: str, *, wait: bool = False) -> dict[str, Any]:
            assert wait and Path(file_path).is_file()
            return {"doc_id": f"synthetic-{document_id}"}

        def get_tree(
            self, doc_id: str, node_summary: bool = False, include_text: bool = True
        ) -> dict[str, Any]:
            return {"doc_id": doc_id}

        def get_page_content(self, doc_id: str, pages: str) -> list[dict[str, Any]]:
            assert doc_id == f"synthetic-{document_id}" and pages == "1"
            return [{"page": 1, "text": quote}]

    pageindex = PageIndexAdapter(SyntheticPageIndex())
    indexer = DocumentIndexWorker(
        DocumentIndexJobStore(TEST_DATABASE_URL),
        tenant,
        "review-fixture-indexer",
        object_store,
        pageindex,
    )
    assert indexer.run_once() == "completed"
    index_configuration = pageindex.configuration_snapshot
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        indexed_configuration = connection.execute(
            "SELECT pageindex_configuration, pageindex_configuration_sha256 "
            "FROM document_indexes WHERE tenant_id = %s AND document_version_id = %s",
            (tenant, document_id),
        ).fetchone()
        assert indexed_configuration == (
            index_configuration["configuration"],
            index_configuration["configuration_sha256"],
        )

    preview = client.post(
        f"/api/v1/cases/{case_id}/source-previews",
        json={"purpose": "contract_review", "task_requires_amount": True},
    )
    assert preview.status_code == 201, preview.text
    frozen = client.post(
        f"/api/v1/cases/{case_id}/input-snapshots",
        json={
            "purpose": "contract_review",
            "task_requires_amount": True,
            "preview_id": preview.json()["preview_id"],
            "expected_case_version": preview.json()["case_version"],
            "manifest_sha256": preview.json()["manifest_sha256"],
        },
    )
    assert frozen.status_code == 200, frozen.text
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        frozen_index_configuration = connection.execute(
            "SELECT pageindex_configuration_sha256 FROM input_snapshot_documents "
            "WHERE tenant_id = %s AND snapshot_id = %s AND document_version_id = %s",
            (tenant, frozen.json()["input_snapshot_id"], document_id),
        ).fetchone()
        assert frozen_index_configuration == (index_configuration["configuration_sha256"],)
        frozen_parse_identity = connection.execute(
            "SELECT parser_version, parsed_sha256 FROM input_snapshot_documents "
            "WHERE tenant_id = %s AND snapshot_id = %s AND document_version_id = %s",
            (tenant, frozen.json()["input_snapshot_id"], document_id),
        ).fetchone()
        document_parse_identity = connection.execute(
            "SELECT parser_version, parsed_sha256 FROM document_versions "
            "WHERE tenant_id = %s AND id = %s",
            (tenant, document_id),
        ).fetchone()
        assert frozen_parse_identity == document_parse_identity
    started = client.post(
        f"/api/v1/cases/{case_id}/runs",
        json={"input_snapshot_id": frozen.json()["input_snapshot_id"]},
        headers={"Idempotency-Key": "review-worker-run-001"},
    )
    assert started.status_code == 202, started.text
    rule_bytes = (
        Path(__file__).parents[1]
        / "docs/development/06-rules/rules/legal.payment_conditions.v1.yaml"
    ).read_bytes()
    skill_bytes = (
        Path(__file__).parents[1] / "skills/cn-procurement-contract-audit/SKILL.md"
    ).read_bytes()
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        snapshot_row = connection.execute(
            "SELECT snapshot FROM run_definition_snapshots WHERE tenant_id = %s AND run_id = %s",
            (tenant, started.json()["run_id"]),
        ).fetchone()
        assert snapshot_row is not None
        frozen_definitions = snapshot_row[0]
        assert frozen_definitions["rule_source"].encode() == rule_bytes
        assert frozen_definitions["dependency_rule_sources"] == {
            "legal.price.definition": (
                Path(__file__).parents[1]
                / "docs/development/06-rules/rules/legal.price.definition.yaml"
            ).read_text(encoding="utf-8")
        }
        assert frozen_definitions["skill_source"].startswith("<!-- development-only:")
        assert frozen_definitions["skill_source"].encode() != skill_bytes
        assert (
            frozen_definitions["manifest"]["skill"]["source_sha256"]
            == hashlib.sha256(frozen_definitions["skill_source"].encode()).hexdigest()
        )
        assert (
            frozen_definitions["manifest"]["skill"]["generator_version"]
            == "development-skill-renderer-v1"
        )
    from contract_app.workflow.definitions import FrozenDefinitionSnapshot
    from contract_app.workflow.model_profile import capture_model_profile

    changed_root = tmp_path / "changed-definitions"
    changed_rule = changed_root / "docs/development/06-rules/rules/legal.payment_conditions.v1.yaml"
    changed_skill = changed_root / "skills/cn-procurement-contract-audit/SKILL.md"
    changed_rule.parent.mkdir(parents=True)
    changed_skill.parent.mkdir(parents=True)
    changed_rule.write_bytes(rule_bytes)
    changed_skill.write_bytes(skill_bytes + b"\n<!-- changed after first request -->\n")
    price_rule_bytes = (
        Path(__file__).parents[1] / "docs/development/06-rules/rules/legal.price.definition.yaml"
    ).read_bytes()
    price_rule_path = changed_rule.parent / "legal.price.definition.yaml"
    price_rule_path.write_bytes(price_rule_bytes)
    changed_definitions = FrozenDefinitionSnapshot.capture(
        changed_root, model_profile=capture_model_profile()
    )
    replayed_run_id = JobStore(TEST_DATABASE_URL).start_bound_run(
        tenant,
        case_id,
        actor,
        UUID(frozen.json()["input_snapshot_id"]),
        "review-worker-run-001",
        changed_definitions.storage_value(),
    )
    assert str(replayed_run_id) == started.json()["run_id"]
    with psycopg.connect(TEST_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        grants = connection.execute(
            "SELECT has_table_privilege(current_user, 'run_definition_snapshots', 'UPDATE'), "
            "has_table_privilege(current_user, 'run_definition_snapshots', 'DELETE')"
        ).fetchone()
        assert grants == (False, False)

    class CandidateModel(ChatModelBase):
        def __init__(self) -> None:
            super().__init__(
                credential=OpenAICredential(api_key=SecretStr("offline-mock")),
                model="offline-mock",
                parameters=ChatModelBase.Parameters(),
                stream=False,
                max_retries=0,
            )
            self.formatter = OpenAIChatFormatter()
            self.calls = 0
            self.injected_conclusion_once = False

        async def _call_api(
            self,
            model_name: str,
            messages: list[Any],
            tools: list[dict[str, Any]] | None = None,
            tool_choice: Any = None,
            **kwargs: Any,
        ) -> Any:
            del model_name, messages, tool_choice, kwargs
            self.calls += 1
            assert {item["function"]["name"] for item in tools or []} == {"locate_clause"}
            if self.calls % 2 == 1:
                return ChatResponse(
                    content=[
                        ToolCallBlock(
                            id="locate-payment", name="locate_clause", input='{"query":"CNY 1000"}'
                        )
                    ],
                    is_last=True,
                    usage=ChatUsage(
                        input_tokens=100 + self.calls, output_tokens=self.calls, time=0.01
                    ),
                )
            if not self.injected_conclusion_once:
                self.injected_conclusion_once = True
                return ChatResponse(
                    content=[TextBlock(text=json.dumps({"judgment": "passed"}))],
                    is_last=True,
                    usage=ChatUsage(
                        input_tokens=100 + self.calls, output_tokens=self.calls, time=0.01
                    ),
                )
            return ChatResponse(
                content=[
                    TextBlock(
                        text=json.dumps(
                            {
                                "observations": [
                                    {
                                        "document_version_id": document_id,
                                        "page": 1,
                                        "quote": quote,
                                    }
                                ],
                                "fact_candidates": [
                                    {
                                        "predicate": "contract.contract_type",
                                        "state": "candidate",
                                        "value": "service",
                                        "evidence": [
                                            {
                                                "document_version_id": document_id,
                                                "page": 1,
                                                "quote": quote,
                                            }
                                        ],
                                    },
                                    {
                                        "predicate": "contract.party_role",
                                        "state": "candidate",
                                        "value": "customer",
                                        "evidence": [
                                            {
                                                "document_version_id": document_id,
                                                "page": 1,
                                                "quote": quote,
                                            }
                                        ],
                                    },
                                    {
                                        "predicate": "contract.jurisdiction",
                                        "state": "candidate",
                                        "value": "CN-mainland",
                                        "evidence": [
                                            {
                                                "document_version_id": document_id,
                                                "page": 1,
                                                "quote": quote,
                                            }
                                        ],
                                    },
                                    {
                                        "predicate": "contract.total_amount",
                                        "state": "candidate",
                                        "value": "CNY 1000",
                                        "evidence": [
                                            {
                                                "document_version_id": document_id,
                                                "page": 1,
                                                "quote": quote,
                                            }
                                        ],
                                    },
                                ],
                                "slot_candidates": [
                                    {
                                        "slot_id": "payment_clause_present",
                                        "fact_type": None,
                                        "state": "filled",
                                        "value": True,
                                        "evidence": [
                                            {
                                                "document_version_id": document_id,
                                                "page": 1,
                                                "quote": quote,
                                            }
                                        ],
                                    },
                                    {
                                        "slot_id": "payment_schedule",
                                        "fact_type": "contract.payment_schedule",
                                        "state": "filled",
                                        "value": [
                                            {
                                                "trigger": "验收",
                                                "amount_type": "percentage",
                                                "percentage": "100",
                                            }
                                        ],
                                        "evidence": [
                                            {
                                                "document_version_id": document_id,
                                                "page": 1,
                                                "quote": quote,
                                            }
                                        ],
                                    },
                                    {
                                        "slot_id": "price_basis",
                                        "fact_type": "contract.price_basis",
                                        "state": "filled",
                                        "value": {
                                            "tax_inclusive_total": "1000",
                                            "tax_exclusive_total": "900",
                                            "tax_amount": tax_amount,
                                            "tax_rate": None,
                                            "currency": "CNY",
                                        },
                                        "evidence": [
                                            {
                                                "document_version_id": document_id,
                                                "page": 1,
                                                "quote": quote,
                                            }
                                        ],
                                    },
                                    {
                                        "slot_id": "price_sources",
                                        "fact_type": "contract.price_source",
                                        "state": "filled",
                                        "value": price_source_candidates,
                                        "evidence": price_source_evidence,
                                    },
                                    {
                                        "slot_id": "payment_basis",
                                        "fact_type": "contract.payment_basis",
                                        "state": "filled",
                                        "value": [
                                            {
                                                "stage": "验收",
                                                "denominator": payment_denominator,
                                                "percentage": "100",
                                                "fixed_amount": fixed_payment_amount,
                                                "residual": False,
                                                "currency": "CNY",
                                            }
                                        ],
                                        "evidence": [
                                            {
                                                "document_version_id": document_id,
                                                "page": 1,
                                                "quote": quote,
                                            }
                                        ],
                                    },
                                    {
                                        "slot_id": "payment_triggers",
                                        "fact_type": "contract.payment_trigger",
                                        "state": "filled",
                                        "value": [
                                            {
                                                "stage": "尾款",
                                                "conditions": ["验收合格"],
                                                "acceptance_required": True,
                                            }
                                        ],
                                        "evidence": [
                                            {
                                                "document_version_id": document_id,
                                                "page": 1,
                                                "quote": quote,
                                            }
                                        ],
                                    },
                                    {
                                        "slot_id": "prepayment_ratio",
                                        "fact_type": None,
                                        "state": "human_review",
                                        "value": None,
                                        "evidence": [],
                                    },
                                    {
                                        "slot_id": "payment_risk_judgment",
                                        "fact_type": None,
                                        "state": "human_review",
                                        "value": None,
                                        "evidence": [],
                                    },
                                ],
                                "limitations": [],
                            },
                            ensure_ascii=False,
                        )
                    )
                ],
                is_last=True,
                usage=ChatUsage(input_tokens=100 + self.calls, output_tokens=self.calls, time=0.01),
            )

    processor = ContractReviewProcessor(
        TEST_DATABASE_URL,
        tenant,
        object_store,
        pageindex,
        CandidateModel(),
        timeout_seconds=5,
        model_profile=__import__(
            "contract_app.workflow.model_profile", fromlist=["capture_model_profile"]
        ).capture_model_profile(),
    )
    processor_errors: list[Exception] = []

    def capture_processor_error(
        payload: dict[str, Any], task: Any, worker_context: Any
    ) -> dict[str, Any]:
        try:
            return processor(payload, task, worker_context)
        except Exception as error:
            processor_errors.append(error)
            raise

    worker = Worker(
        JobStore(TEST_DATABASE_URL),
        tenant,
        "review-fixture-worker",
        {},
        contextual_handlers={"contract_review": capture_processor_error},
        lease_seconds=15,
        max_attempts=2,
        retry_delay_seconds=0,
    )
    assert worker.run_once() == "retried"
    pending_results = client.get(f"/api/v1/runs/{started.json()['run_id']}/results")
    assert pending_results.status_code == 200, pending_results.text
    assert pending_results.json()["results"] == []
    assert pending_results.json()["reported_model_usage"] == {
        "model_requests": 2,
        "model_responses": 2,
        "usage_reports": 2,
        "input_tokens": 203,
        "output_tokens": 3,
        "complete": True,
    }
    with psycopg.connect(TEST_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        retry_state = connection.execute(
            "SELECT status, attempt, last_error FROM tasks WHERE run_id = %s",
            (started.json()["run_id"],),
        ).fetchone()
        assert retry_state == ("queued", 1, "ValueError")

    worker_outcome = worker.run_once()
    assert worker_outcome == "completed", repr(processor_errors)
    results = client.get(f"/api/v1/runs/{started.json()['run_id']}/results")
    assert results.status_code == 200, results.text
    reported_usage = results.json()["reported_model_usage"]
    assert reported_usage == {
        "model_requests": 4,
        "model_responses": 4,
        "usage_reports": 4,
        "input_tokens": 410,
        "output_tokens": 10,
        "complete": True,
    }
    assert len(results.json()["results"]) == 2
    result = next(
        item for item in results.json()["results"] if item["rule_id"] == "legal.payment_conditions"
    )
    price_result = next(
        item for item in results.json()["results"] if item["rule_id"] == "legal.price.definition"
    )
    assert result["judgment"] == "undetermined"
    assert price_result["judgment"] == "undetermined"
    assert result["severity"] == "undetermined"
    assert len(result["dependency_result_references"]) == 1
    dependency_reference = result["dependency_result_references"][0]
    assert dependency_reference["check_id"] == "legal.price.definition"
    assert dependency_reference["result_id"] == price_result["result_id"]
    assert dependency_reference["source_run_id"] == started.json()["run_id"]
    assert dependency_reference["is_completion_proof"] is False
    assert result["consumed_dependency_result_references"] == []
    assert (
        dependency_reference["assessment_status"]
        == (price_result["rule_plan_execution"]["assessment_status"])
    )
    assert result["rule_plan_execution"]["blocked_by_checks"] == ["legal.price.definition"]
    assert price_result["fact_set_snapshot_id"] == result["fact_set_snapshot_id"]
    assert (
        price_result["rule_plan_execution"]["input_snapshot_id"]
        == (results.json()["run"]["input_snapshot_id"])
    )
    assert price_result["rule_plan_execution"]["is_dependency_completion_proof"] is False
    rule_plan_execution = result["rule_plan_execution"]
    with psycopg.connect(TEST_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        stored_results = connection.execute(
            "SELECT result, content_hash FROM task_results WHERE run_id = %s",
            (started.json()["run_id"],),
        ).fetchall()
    assert len(stored_results) == 1
    assert result["task_result_sha256"] == stored_results[0][1]
    assert price_result["task_result_sha256"] == stored_results[0][1]
    persisted_by_id = {item["result_id"]: item for item in stored_results[0][0]["results"]}
    assert set(persisted_by_id) == {item["result_id"] for item in results.json()["results"]}
    assert all(
        persisted_by_id[item["result_id"]]["slot_candidates"] == item["slot_candidates"]
        for item in results.json()["results"]
    )
    price_tasks = client.get(f"/api/v1/human-tasks?run_id={started.json()['run_id']}&status=open")
    assert price_tasks.status_code == 200, price_tasks.text
    price_task = next(
        item for item in price_tasks.json()["items"] if item["type"] == "price_review"
    )
    decision = client.post(
        f"/api/v1/human-tasks/{price_task['task_id']}/decisions",
        headers={"Idempotency-Key": "price-review-decision-001"},
        json={
            "version": price_task["version"],
            "decision": "record_review",
            "reason": "已核对引用；价款结论仍待定。",
        },
    )
    assert decision.status_code == 200, decision.text
    assert decision.json()["status"] == "completed"
    assert decision.json()["decision"] == "record_review"
    replay = client.post(
        f"/api/v1/human-tasks/{price_task['task_id']}/decisions",
        headers={"Idempotency-Key": "price-review-decision-001"},
        json={
            "version": price_task["version"],
            "decision": "record_review",
            "reason": "已核对引用；价款结论仍待定。",
        },
    )
    assert replay.status_code == 200, replay.text
    assert replay.json() == decision.json()
    wrong_decision = client.post(
        f"/api/v1/human-tasks/{price_task['task_id']}/decisions",
        headers={"Idempotency-Key": "price-review-fact-001"},
        json={
            "version": price_task["version"],
            "decision": "confirm_fact",
            "reason": "不应允许价款任务变成事实确认。",
        },
    )
    assert wrong_decision.status_code == 404
    unchanged_results = client.get(f"/api/v1/runs/{started.json()['run_id']}/results")
    assert unchanged_results.status_code == 200, unchanged_results.text
    unchanged_price = next(
        item
        for item in unchanged_results.json()["results"]
        if item["rule_id"] == "legal.price.definition"
    )
    assert unchanged_price["judgment"] == "undetermined"
    assert next(
        item
        for item in unchanged_results.json()["results"]
        if item["rule_id"] == "legal.payment_conditions"
    )["rule_plan_execution"]["blocked_by_checks"] == ["legal.price.definition"]
    transferred = client.post(
        f"/api/v1/cases/{case_id}/runs",
        json={
            "input_snapshot_id": frozen.json()["input_snapshot_id"],
            "dependency_source_run_id": started.json()["run_id"],
        },
        headers={"Idempotency-Key": "review-worker-run-transfer-001"},
    )
    assert transferred.status_code == 202, transferred.text
    with psycopg.connect(TEST_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        transferred_payload = connection.execute(
            "SELECT payload FROM tasks WHERE run_id = %s AND task_type = 'contract_review'",
            (transferred.json()["run_id"],),
        ).fetchone()
    assert transferred_payload is not None
    transferred_references = transferred_payload[0]["dependency_result_references"]
    assert len(transferred_references) == 1
    assert transferred_references[0]["result_id"] == price_result["result_id"]
    assert transferred_references[0]["source_run_id"] == started.json()["run_id"]
    assert transferred_references[0]["is_completion_proof"] is False
    assert worker.run_once() == "completed", repr(processor_errors)
    transferred_results = client.get(f"/api/v1/runs/{transferred.json()['run_id']}/results")
    assert transferred_results.status_code == 200, transferred_results.text
    transferred_payment_result = next(
        item
        for item in transferred_results.json()["results"]
        if item["rule_id"] == "legal.payment_conditions"
    )
    assert transferred_payment_result["judgment"] == "undetermined"
    assert transferred_payment_result["rule_plan_execution"]["blocked_by_checks"] == [
        "legal.price.definition"
    ]
    consumed_references = transferred_payment_result["consumed_dependency_result_references"]
    assert len(consumed_references) == 1
    assert consumed_references[0]["result_id"] == price_result["result_id"]
    assert consumed_references[0]["source_run_id"] == started.json()["run_id"]
    assert consumed_references[0]["is_completion_proof"] is False
    from contract_app.workflow.dependency_result_reference import (
        resolve_dependency_result_reference,
    )

    price_rule_bytes = (
        Path(__file__).parents[1] / "docs/development/06-rules/rules/legal.price.definition.yaml"
    ).read_bytes()
    price_reference = resolve_dependency_result_reference(
        JobStore(TEST_DATABASE_URL).get_run_results(tenant, actor, UUID(started.json()["run_id"])),
        expected_case_id=case_id,
        expected_input_snapshot_id=UUID(frozen.json()["input_snapshot_id"]),
        expected_fact_set_snapshot_id=(
            UUID(price_result["fact_set_snapshot_id"])
            if price_result["fact_set_snapshot_id"]
            else None
        ),
        check_id="legal.price.definition",
        dependency_rule_id="legal.price.definition",
        dependency_rule_version="1.0.0",
        dependency_rule_sha256=hashlib.sha256(price_rule_bytes).hexdigest(),
    )
    assert price_reference.result_id == UUID(price_result["result_id"])
    assert price_reference.task_result_sha256 == stored_results[0][1]
    assert price_reference.evidence_manifest_sha256 != "0" * 64
    assert rule_plan_execution["kind"] == "rule_plan_assessment"
    assert rule_plan_execution["input_snapshot_id"] == results.json()["run"]["input_snapshot_id"]
    assert rule_plan_execution["fact_set_snapshot_id"] == result["fact_set_snapshot_id"]
    assert rule_plan_execution["assessment_status"] == "undetermined"
    assert rule_plan_execution["blocked_by_checks"] == ["legal.price.definition"]
    assert rule_plan_execution["is_dependency_completion_proof"] is False
    assert len(rule_plan_execution["assessment_sha256"]) == 64
    limitations_text = "".join(result["limitations"])
    assert "自动通过" in limitations_text
    assert "规则适用性尚未确定" in limitations_text
    assert "legal.price.definition" in limitations_text
    assert "发布阻断" in limitations_text
    expected_price_hint = {
        "normal": "含税价减未税价与税额候选一致",
        "tax_conflict": "含税价、未税价与税额不一致",
        "amount_conflict": "价款来源不一致",
        "missing_denominator": "付款比例分母未明确",
    }[price_case]
    assert expected_price_hint in limitations_text
    if price_case == "amount_conflict":
        assert "固定付款金额与付款比例计算不一致" in limitations_text
    elif price_case != "missing_denominator":
        assert "固定付款金额与按合同总价候选计算的比例金额一致" in limitations_text
    execution_trace = result["execution_trace"]
    assert execution_trace["schema_version"] == 1
    trace_event_types = {event["event_type"] for event in execution_trace["events"]}
    assert {"model_request", "model_response", "tool_exchange"} <= trace_event_types
    assert execution_trace["model_usage"] == {
        "model_responses": 2,
        "usage_reports": 2,
        "input_tokens": 207,
        "output_tokens": 7,
        "complete": True,
    }
    assert quote not in json.dumps(execution_trace, ensure_ascii=False)
    assert "Payment" not in json.dumps(execution_trace, ensure_ascii=False)
    manifest = result["definition_manifest"]
    rule_bytes = (
        Path(__file__).parents[1]
        / "docs/development/06-rules/rules/legal.payment_conditions.v1.yaml"
    ).read_bytes()
    skill_bytes = (
        Path(__file__).parents[1] / "skills/cn-procurement-contract-audit/SKILL.md"
    ).read_bytes()
    assert manifest["kind"] == "development_execution_definitions"
    assert manifest["rule"]["source_sha256"] == hashlib.sha256(rule_bytes).hexdigest()
    assert (
        manifest["skill"]["source_sha256"]
        == hashlib.sha256(frozen_definitions["skill_source"].encode()).hexdigest()
    )
    assert manifest["skill"]["status"] == "development_candidate"
    assert manifest["tool_contract"] == frozen_definitions["manifest"]["tool_contract"]
    manifest_for_hash = {key: value for key, value in manifest.items() if key != "manifest_sha256"}
    assert (
        manifest["manifest_sha256"]
        == hashlib.sha256(
            json.dumps(manifest_for_hash, sort_keys=True, separators=(",", ":")).encode()
        ).hexdigest()
    )
    assert result["evidence"][0]["quote"] == quote
    assert result["evidence"][0]["document_version_id"] == document_id
    assert [item["predicate"] for item in result["fact_candidates"]] == [
        "contract.contract_type",
        "contract.party_role",
        "contract.jurisdiction",
        "contract.total_amount",
    ]
    assert all(
        item["evidence_ids"] == [result["evidence"][0]["evidence_id"]]
        for item in result["fact_candidates"]
    )
    slot_candidates = {item["slot_id"]: item for item in result["slot_candidates"]}
    price_slot_candidates = {item["slot_id"]: item for item in price_result["slot_candidates"]}
    assert slot_candidates["payment_schedule"]["state"] == "filled"
    assert slot_candidates["payment_schedule"]["fact_type"] == "contract.payment_schedule"
    assert slot_candidates["payment_schedule"]["value"][0]["percentage"] == "100"
    assert price_slot_candidates["price_basis"]["value"] == {
        "tax_inclusive_total": "1000",
        "tax_exclusive_total": "900",
        "tax_amount": tax_amount,
        "tax_rate": None,
        "currency": "CNY",
    }
    assert price_slot_candidates["price_sources"]["value"][0]["source_type"] == "contract_body"
    assert [item["amount"] for item in price_slot_candidates["price_sources"]["value"]] == (
        ["1000", "1200"] if price_case == "amount_conflict" else ["1000"]
    )
    assert price_slot_candidates["payment_basis"]["value"][0]["denominator"] == (
        payment_denominator
    )
    source_evidence_by_quote = {
        item["quote"]: item["evidence_id"] for item in price_result["evidence"]
    }
    expected_source_evidence_quotes = [item["quote"] for item in price_source_evidence]
    assert price_slot_candidates["price_sources"]["evidence_ids"] == [
        source_evidence_by_quote[quote_text] for quote_text in expected_source_evidence_quotes
    ]
    assert all(
        price_slot_candidates[slot_id]["evidence_ids"]
        for slot_id in ("price_basis", "payment_basis")
    )
    price_evidence_ids = {item["evidence_id"] for item in price_result["evidence"]}
    assert all(
        set(price_slot_candidates[slot_id]["evidence_ids"]).issubset(price_evidence_ids)
        for slot_id in ("price_basis", "price_sources", "payment_basis")
    )
    assert all(
        price_slot_candidates[slot_id]["state"] == "filled"
        for slot_id in ("price_basis", "price_sources", "payment_basis")
    )
    assert slot_candidates["prepayment_ratio"] == {
        "slot_id": "prepayment_ratio",
        "state": "human_review",
        "value": None,
        "fact_type": None,
        "evidence_ids": [],
    }
    assert slot_candidates["payment_risk_judgment"]["value"] is None

    dependency_run_id = JobStore(TEST_DATABASE_URL).start_bound_run(
        tenant,
        case_id,
        actor,
        UUID(frozen.json()["input_snapshot_id"]),
        "review-worker-run-dependency-source-001",
        changed_definitions.storage_value(),
    )
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        dependency_snapshot = connection.execute(
            "SELECT snapshot FROM run_definition_snapshots WHERE tenant_id = %s AND run_id = %s",
            (tenant, dependency_run_id),
        ).fetchone()
    assert dependency_snapshot is not None
    assert dependency_snapshot[0]["dependency_rule_sources"] == {
        "legal.price.definition": price_rule_bytes.decode("utf-8")
    }
    price_rule_path.write_text("changed after queueing", encoding="utf-8")

    assert worker.run_once() == "completed", repr(processor_errors)
    dependency_results = client.get(f"/api/v1/runs/{dependency_run_id}/results")
    assert dependency_results.status_code == 200, dependency_results.text
    dependency_result = dependency_results.json()["results"][0]
    assert dependency_result["definition_manifest"]["dependency_graph"]["status"] == (
        "graph_resolved"
    )
    assert dependency_result["rule_plan_execution"]["blocked_by_checks"] == [
        "legal.price.definition"
    ]
    assert dependency_result["judgment"] == "undetermined"

    confirmation = client.post(
        f"/api/v1/runs/{started.json()['run_id']}/fact-set-snapshots",
        headers={"Idempotency-Key": "confirm-facts-run-0001"},
        json={
            "confirmations": [
                {
                    "predicate": item["predicate"],
                    "value": item["value"],
                    "evidence_ids": item["evidence_ids"],
                }
                for item in result["fact_candidates"]
            ],
            "reason": "已核对合同类型、签约角色与适用法条。",
        },
    )
    assert confirmation.status_code == 201, confirmation.text
    fact_set_snapshot = confirmation.json()
    assert fact_set_snapshot["version"] == 1
    assert fact_set_snapshot["input_snapshot_id"] == frozen.json()["input_snapshot_id"]
    assert {key: value["value"] for key, value in fact_set_snapshot["facts"].items()} == {
        "contract.contract_type": "service",
        "contract.party_role": "customer",
        "contract.jurisdiction": "CN-mainland",
        "contract.total_amount": "CNY 1000",
    }
    assert all(
        value["verification_status"] == "verified"
        and value["confirmed_by"] == str(actor)
        and value["evidence_ids"] == [result["evidence"][0]["evidence_id"]]
        for value in fact_set_snapshot["facts"].values()
    )
    unauthorized_client = _client(tenant, uuid4(), object_store)
    unauthorized_confirmation = unauthorized_client.post(
        f"/api/v1/runs/{started.json()['run_id']}/fact-set-snapshots",
        headers={"Idempotency-Key": "confirm-facts-denied-01"},
        json={
            "confirmations": [
                {
                    "predicate": result["fact_candidates"][0]["predicate"],
                    "value": result["fact_candidates"][0]["value"],
                    "evidence_ids": result["fact_candidates"][0]["evidence_ids"],
                }
            ],
            "reason": "Unauthorized actor attempt.",
        },
    )
    assert unauthorized_confirmation.status_code == 404
    with psycopg.connect(TEST_DATABASE_URL) as connection:
        append_only_privileges = connection.execute(
            "SELECT has_table_privilege(current_user, 'fact_set_snapshots', 'UPDATE'), "
            "has_table_privilege(current_user, 'fact_set_snapshots', 'DELETE')"
        ).fetchone()
    assert append_only_privileges == (False, False)
    replayed_confirmation = client.post(
        f"/api/v1/runs/{started.json()['run_id']}/fact-set-snapshots",
        headers={"Idempotency-Key": "confirm-facts-run-0001"},
        json={
            "confirmations": [
                {
                    "predicate": item["predicate"],
                    "value": item["value"],
                    "evidence_ids": item["evidence_ids"],
                }
                for item in result["fact_candidates"]
            ],
            "reason": "已核对合同类型、签约角色与适用法条。",
        },
    )
    assert replayed_confirmation.status_code == 201
    assert (
        replayed_confirmation.json()["fact_set_snapshot_id"]
        == fact_set_snapshot["fact_set_snapshot_id"]
    )

    followup_run = client.post(
        f"/api/v1/cases/{case_id}/runs",
        headers={"Idempotency-Key": "review-worker-run-confirmed-001"},
        json={"input_snapshot_id": frozen.json()["input_snapshot_id"]},
    )
    assert followup_run.status_code == 202, followup_run.text
    replayed_followup_run = client.post(
        f"/api/v1/cases/{case_id}/runs",
        headers={"Idempotency-Key": "review-worker-run-confirmed-001"},
        json={"input_snapshot_id": frozen.json()["input_snapshot_id"]},
    )
    assert replayed_followup_run.status_code == 202
    assert replayed_followup_run.json()["run_id"] == followup_run.json()["run_id"]
    with psycopg.connect(TEST_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        followup_payload = connection.execute(
            "SELECT payload FROM tasks WHERE run_id = %s", (followup_run.json()["run_id"],)
        ).fetchone()[0]
    assert followup_payload["fact_set_snapshot_id"] == fact_set_snapshot["fact_set_snapshot_id"]

    followup_errors: list[Exception] = []

    def capture_followup_error(
        payload: dict[str, Any], task: Any, worker_context: Any
    ) -> dict[str, Any]:
        try:
            return processor(payload, task, worker_context)
        except Exception as error:
            followup_errors.append(error)
            raise

    followup_worker = Worker(
        JobStore(TEST_DATABASE_URL),
        tenant,
        "fact-bound-worker",
        {},
        contextual_handlers={"contract_review": capture_followup_error},
        lease_seconds=15,
        max_attempts=1,
        retry_delay_seconds=0,
    )
    assert followup_worker.run_once() == "completed", repr(followup_errors)
    followup_results = client.get(f"/api/v1/runs/{followup_run.json()['run_id']}/results")
    assert followup_results.status_code == 200, followup_results.text
    followup_result = followup_results.json()["results"][0]
    assert followup_result["fact_set_snapshot_id"] == fact_set_snapshot["fact_set_snapshot_id"]
    assert len(followup_result["task_result_sha256"]) == 64
    followup_plan_execution = followup_result["rule_plan_execution"]
    confirmed_fact_values = {
        predicate: item["value"] for predicate, item in fact_set_snapshot["facts"].items()
    }
    expected_facts_sha256 = hashlib.sha256(
        json.dumps(
            confirmed_fact_values,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()
    assert followup_plan_execution["input_snapshot_id"] == frozen.json()["input_snapshot_id"]
    assert (
        followup_plan_execution["fact_set_snapshot_id"] == fact_set_snapshot["fact_set_snapshot_id"]
    )
    assert followup_plan_execution["facts_sha256"] == expected_facts_sha256
    assert followup_plan_execution["assessment_status"] == "blocked"
    assert followup_plan_execution["blocked_by_checks"] == ["legal.price.definition"]
    assert followup_plan_execution["is_dependency_completion_proof"] is False
    followup_limitations = "".join(followup_result["limitations"])
    assert followup_result["judgment"] == "undetermined"
    assert "规则适用性尚未确定" not in followup_limitations
    assert "前置检查尚未完成" in followup_limitations
    assert "legal.price.definition" in followup_limitations

    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        restricted = connection.execute(
            "UPDATE document_grants SET allow_amount = false "
            "WHERE tenant_id = %s AND actor_id = %s AND document_version_id = %s "
            "AND purpose = 'contract_review' AND revoked_at IS NULL",
            (tenant, actor, document_id),
        )
        assert restricted.rowcount == 1
    replay_after_amount_revocation = client.post(
        f"/api/v1/runs/{started.json()['run_id']}/fact-set-snapshots",
        headers={"Idempotency-Key": "confirm-facts-run-0001"},
        json={
            "confirmations": [
                {
                    "predicate": item["predicate"],
                    "value": item["value"],
                    "evidence_ids": item["evidence_ids"],
                }
                for item in result["fact_candidates"]
            ],
            "reason": "已核对合同类型、签约角色与适用法条。",
        },
    )
    assert replay_after_amount_revocation.status_code == 404
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        connection.execute(
            "UPDATE document_grants SET allow_amount = true "
            "WHERE tenant_id = %s AND actor_id = %s AND document_version_id = %s "
            "AND purpose = 'contract_review'",
            (tenant, actor, document_id),
        )

    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        revoked = connection.execute(
            "UPDATE document_grants SET revoked_at = now() "
            "WHERE tenant_id = %s AND actor_id = %s AND document_version_id = %s "
            "AND purpose = 'contract_review' AND revoked_at IS NULL",
            (tenant, actor, document_id),
        )
        assert revoked.rowcount == 1
    replay_after_revocation = client.post(
        f"/api/v1/runs/{started.json()['run_id']}/fact-set-snapshots",
        headers={"Idempotency-Key": "confirm-facts-run-0001"},
        json={
            "confirmations": [
                {
                    "predicate": item["predicate"],
                    "value": item["value"],
                    "evidence_ids": item["evidence_ids"],
                }
                for item in result["fact_candidates"]
            ],
            "reason": "已核对合同类型、签约角色与适用法条。",
        },
    )
    assert replay_after_revocation.status_code == 404
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        connection.execute(
            "UPDATE document_grants SET revoked_at = NULL "
            "WHERE tenant_id = %s AND actor_id = %s AND document_version_id = %s "
            "AND purpose = 'contract_review'",
            (tenant, actor, document_id),
        )

    # Distinct idempotency keys racing on the same input version must not
    # create duplicate version numbers or both report success.
    concurrent_confirmation = {
        "confirmations": [
            {
                "predicate": "contract.contract_type",
                "value": "service",
                "evidence_ids": result["fact_candidates"][0]["evidence_ids"],
            }
        ],
        "reason": "并发版本冲突验证。",
    }
    start_barrier = Barrier(2)

    def submit_concurrent_confirmation(key: str) -> int:
        with _client(tenant, actor, object_store) as concurrent_client:
            start_barrier.wait(timeout=10)
            response = concurrent_client.post(
                f"/api/v1/runs/{started.json()['run_id']}/fact-set-snapshots",
                headers={"Idempotency-Key": key},
                json=concurrent_confirmation,
            )
            return response.status_code

    with ThreadPoolExecutor(max_workers=2) as executor:
        concurrent_statuses = list(
            executor.map(
                submit_concurrent_confirmation,
                ("confirm-race-key-0001", "confirm-race-key-0002"),
            )
        )
    assert sorted(concurrent_statuses) == [201, 409]
    with psycopg.connect(TEST_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        concurrent_versions = connection.execute(
            "SELECT array_agg(version ORDER BY version) FROM fact_set_snapshots "
            "WHERE tenant_id = %s AND case_id = %s AND input_snapshot_id = %s",
            (tenant, case_id, frozen.json()["input_snapshot_id"]),
        ).fetchone()[0]
    assert concurrent_versions == [1, 2]

    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        stored = connection.execute(
            "SELECT task_id, result FROM task_results WHERE tenant_id = %s AND run_id = %s",
            (tenant, started.json()["run_id"]),
        ).fetchone()
        assert stored is not None
        source_record_task_id, result_with_external_source = stored
        original_result_serialized = json.dumps(
            result_with_external_source,
            sort_keys=True,
            ensure_ascii=False,
            separators=(",", ":"),
        )
        original_result_hash = hashlib.sha256(
            original_result_serialized.encode("utf-8")
        ).hexdigest()
        result_with_external_source["results"][0]["evidence"].append(
            {
                "evidence_id": str(uuid4()),
                "kind": "source_record",
                "support": "context",
                "source_version": "external-source-version",
                "source_record_ref": "finance:restricted-record",
            }
        )
        source_record_serialized = json.dumps(
            result_with_external_source,
            sort_keys=True,
            ensure_ascii=False,
            separators=(",", ":"),
        )
        connection.execute(
            "UPDATE task_results SET result = %s::jsonb, content_hash = %s "
            "WHERE tenant_id = %s AND task_id = %s",
            (
                source_record_serialized,
                hashlib.sha256(source_record_serialized.encode("utf-8")).hexdigest(),
                tenant,
                source_record_task_id,
            ),
        )
    with pytest.raises(RetrievalDenied, match="result_source_not_authorized"):
        JobStore(TEST_DATABASE_URL).get_run_results(tenant, actor, UUID(started.json()["run_id"]))
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        connection.execute(
            "UPDATE task_results SET result = %s::jsonb, content_hash = %s "
            "WHERE tenant_id = %s AND task_id = %s",
            (
                original_result_serialized,
                original_result_hash,
                tenant,
                source_record_task_id,
            ),
        )

    alternate_trace = ExecutionTraceRecorder()
    for index, event in enumerate(execution_trace["events"]):
        alternate_trace.record_digest(
            event["event_type"],
            hashlib.sha256(b"different persisted payload").hexdigest()
            if index == 0
            else event["payload_sha256"],
            event["payload_bytes"],
            input_tokens=event.get("input_tokens"),
            output_tokens=event.get("output_tokens"),
        )
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        stored = connection.execute(
            "SELECT task_id, result FROM task_results WHERE tenant_id = %s AND run_id = %s",
            (tenant, started.json()["run_id"]),
        ).fetchone()
        assert stored is not None
        stored_task_id, mutated_result = stored
        for result_item in mutated_result["results"]:
            result_item["execution_trace"] = alternate_trace.snapshot()
        serialized_result = json.dumps(
            mutated_result, sort_keys=True, ensure_ascii=False, separators=(",", ":")
        )
        connection.execute(
            "UPDATE task_results SET result = %s::jsonb, content_hash = %s "
            "WHERE tenant_id = %s AND task_id = %s",
            (
                serialized_result,
                hashlib.sha256(serialized_result.encode("utf-8")).hexdigest(),
                tenant,
                stored_task_id,
            ),
        )
    with pytest.raises(ValueError, match="stored_execution_trace_ledger_mismatch"):
        JobStore(TEST_DATABASE_URL).get_run_results(tenant, actor, UUID(started.json()["run_id"]))
    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        connection.execute(
            "UPDATE task_results SET content_hash = %s WHERE tenant_id = %s AND run_id = %s",
            ("0" * 64, tenant, started.json()["run_id"]),
        )
    with pytest.raises(ValueError, match="stored_run_result_content_hash_invalid"):
        JobStore(TEST_DATABASE_URL).get_run_results(tenant, actor, UUID(started.json()["run_id"]))


def test_customer_contract_upload_is_separately_authorized_for_upstream_comparison(
    tmp_path: Path,
) -> None:
    assert TEST_DATABASE_URL is not None
    assert TEST_ADMIN_DATABASE_URL is not None
    tenant, case_id, actor = _case()
    client = _client(tenant, actor, LocalObjectStore(tmp_path))
    upload_route = f"/api/v1/cases/{case_id}/documents"
    valid_until = datetime.now(UTC) + timedelta(minutes=30)

    main_upload = client.post(
        upload_route,
        headers={"Idempotency-Key": "upstream-main-001"},
        data={"role": "primary_contract"},
        files={"file": ("purchase.pdf", _pdf("采购合同"), PDF_MEDIA_TYPE)},
    )
    customer_upload = client.post(
        upload_route,
        headers={"Idempotency-Key": "upstream-customer-001"},
        data={"role": "customer_contract"},
        files={"file": ("customer.pdf", _pdf("客户合同"), PDF_MEDIA_TYPE)},
    )
    assert main_upload.status_code == 201, main_upload.text
    assert customer_upload.status_code == 201, customer_upload.text
    assert customer_upload.json()["role"] == "customer_contract"
    assert customer_upload.json()["authorization_status"] == "not_granted_by_upload"

    selection_route = f"/api/v1/cases/{case_id}/document-selections"
    main_document_id = main_upload.json()["document_version_id"]
    customer_document_id = customer_upload.json()["document_version_id"]
    unconfirmed_upstream = client.post(
        selection_route,
        json={
            "document_version_id": customer_document_id,
            "relation": "upstream",
            "purpose": "upstream_comparison",
            "allow_amount": True,
            "valid_until": valid_until.isoformat(),
        },
    )
    assert unconfirmed_upstream.status_code == 422, unconfirmed_upstream.text
    main_review = client.post(
        selection_route,
        json={
            "document_version_id": main_document_id,
            "relation": "main",
            "purpose": "contract_review",
            "allow_amount": True,
            "valid_until": valid_until.isoformat(),
        },
    )
    main_comparison = client.post(
        selection_route,
        json={
            "document_version_id": main_document_id,
            "relation": "main",
            "purpose": "upstream_comparison",
            "allow_amount": True,
            "valid_until": valid_until.isoformat(),
        },
    )
    upstream = client.post(
        selection_route,
        json={
            "document_version_id": customer_document_id,
            "relation": "upstream",
            "purpose": "upstream_comparison",
            "confirm_project_relation": True,
            "allow_amount": True,
            "valid_until": valid_until.isoformat(),
        },
    )
    assert main_review.status_code == 201, main_review.text
    assert main_comparison.status_code == 201, main_comparison.text
    assert upstream.status_code == 201, upstream.text
    assert main_review.json()["purpose"] == "contract_review"
    assert main_comparison.json()["purpose"] == "upstream_comparison"
    assert upstream.json()["relation"] == "upstream"

    with psycopg.connect(TEST_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        customer_relation = connection.execute(
            "SELECT relation, relation_ref, relation_verified_at, selected_by "
            "FROM contract_selections "
            "WHERE tenant_id = %s AND case_id = %s AND document_version_id = %s",
            (tenant, case_id, customer_document_id),
        ).fetchone()
        assert customer_relation[0] == "upstream"
        assert customer_relation[1] == f"case:{case_id}"
        assert customer_relation[2] is not None
        assert customer_relation[3] == actor
        grant_purposes = connection.execute(
            "SELECT document_version_id, purpose, allow_text, allow_amount "
            "FROM document_grants WHERE tenant_id = %s ORDER BY document_version_id, purpose",
            (tenant,),
        ).fetchall()
    assert {(row[0], row[1]) for row in grant_purposes} == {
        (UUID(main_document_id), "contract_review"),
        (UUID(main_document_id), "upstream_comparison"),
        (UUID(customer_document_id), "upstream_comparison"),
    }
    assert all(row[2] and row[3] for row in grant_purposes)

    documents = client.get(f"/api/v1/cases/{case_id}/documents")
    assert documents.status_code == 200, documents.text
    customer_entry = next(
        item for item in documents.json() if item["document_version_id"] == customer_document_id
    )
    assert customer_entry["purpose_authorizations"][0]["purpose"] == "upstream_comparison"
    assert customer_entry["purpose_authorizations"][0]["allow_text"] is True
    assert customer_entry["purpose_authorizations"][0]["status"] == "authorized"

    snapshots = InputSnapshotStore(TEST_DATABASE_URL)
    regular_preview = snapshots.preview(TaskAccess(tenant, case_id, actor, "contract_review"))
    assert {item.relation for item in regular_preview.items} == {"main"}
    comparison_preview = snapshots.preview(
        TaskAccess(tenant, case_id, actor, "upstream_comparison", task_requires_amount=True)
    )
    assert {item.relation for item in comparison_preview.items} == {"main", "upstream"}

    with psycopg.connect(TEST_ADMIN_DATABASE_URL) as connection:
        connection.execute("SELECT set_config('app.tenant_id', %s, true)", (str(tenant),))
        connection.execute(
            "UPDATE document_grants SET valid_until = clock_timestamp() - interval '1 second' "
            "WHERE tenant_id = %s AND actor_id = %s AND document_version_id = %s "
            "AND purpose = 'upstream_comparison' AND revoked_at IS NULL",
            (tenant, actor, customer_document_id),
        )
    expired_documents = client.get(f"/api/v1/cases/{case_id}/documents")
    expired_entry = next(
        item
        for item in expired_documents.json()
        if item["document_version_id"] == customer_document_id
    )
    assert expired_entry["purpose_authorizations"][0]["status"] == "expired"

    renewed = client.post(
        selection_route,
        json={
            "document_version_id": customer_document_id,
            "relation": "upstream",
            "purpose": "upstream_comparison",
            "confirm_project_relation": True,
            "allow_amount": True,
            "valid_until": (datetime.now(UTC) + timedelta(hours=1)).isoformat(),
        },
    )
    assert renewed.status_code == 201, renewed.text
    assert renewed.json()["selection_id"] == upstream.json()["selection_id"]
    assert renewed.json()["case_version"] == upstream.json()["case_version"] + 1

    revoked = client.post(
        f"{selection_route}/{renewed.json()['selection_id']}/revoke",
        headers={"Idempotency-Key": "upstream-revoke-001"},
        json={
            "expected_case_version": renewed.json()["case_version"],
            "reason": "取消本次上下游关联",
        },
    )
    assert revoked.status_code == 200, revoked.text
    assert revoked.json()["revoked_grant_count"] == 1
    comparison_after_revoke = snapshots.preview(
        TaskAccess(tenant, case_id, actor, "upstream_comparison")
    )
    assert {item.relation for item in comparison_after_revoke.items} == {"main"}
